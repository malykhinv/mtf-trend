from __future__ import annotations

import json
from dataclasses import replace
from datetime import datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from bot.data.io.excel_rows import StoredSignalRow, StoredStateRow, StoredTradeRow
from bot.data.io.storage import Storage
from bot.data.repositories.signal_repository import SignalRepository
from bot.data.repositories.state_repository import StateRepository
from bot.data.repositories.trade_repository import TradeRepository
from bot.domain.enums import BreakDirection, Exchange, Side, Timeframe, TradeStatus
from bot.domain.models.entities import Candle, Signal, Thresholds, Trade
from bot.domain.models.metadata import SignalMetadata, TradeMetadata
from bot.domain.services.metrics_service import SelectionMetricsSnapshot


def _create_storage(tmp_path: Path) -> tuple[Storage, Path]:
    path = tmp_path / "storage.xlsx"
    storage = Storage(path)
    return storage, path


def _build_signal(identifier: str, score: float, triggered_at: datetime) -> Signal:
    candle = Candle(
        id=f"candle-{identifier}",
        symbol="BTCUSDT",
        exchange=Exchange.BINANCE,
        timeframe=Timeframe.M15,
        open=100.0,
        high=110.0,
        low=95.0,
        close=105.0,
        volume=1234.0,
        quote_volume=5678.0,
        started_at=triggered_at - timedelta(minutes=15),
        closed_at=triggered_at,
    )
    thresholds = Thresholds(
        id=f"thr-{identifier}",
        min_relative_volume=1.0,
        max_relative_volume=5.0,
        metadata={"timeframe": Timeframe.M15.value},
    )
    base = score / 10.0
    snapshot = SelectionMetricsSnapshot(
        atr=1.0 + base,
        average_volume=2.0 + base,
        momentum=3.0 + base,
        pct_move=4.0 + base,
        relative_volume=5.0 + base,
        atr_multiple=6.0 + base,
        upper_wick_pct=7.0 + base,
        body_pct=8.0 + base,
        lower_wick_pct=9.0 + base,
        pct_to_high=10.0 + base,
        pct_to_low=-(11.0 + base),
        pct_to_high_break=12.0 + base,
        pct_to_low_break=13.0 + base,
        break_direction=BreakDirection.NONE,
    )
    return Signal(
        id=f"sig-{identifier}",
        candle=candle,
        candle_id=candle.id,
        side=Side.LONG,
        direction=BreakDirection.HIGH_FIRST,
        score=score,
        triggered_at=triggered_at,
        thresholds=thresholds,
        timeframe=Timeframe.M15,
        created_at=triggered_at,
        updated_at=triggered_at,
        allow_long=True,
        allow_short=True,
        metrics_snapshot=snapshot,
        metadata=SignalMetadata(extra={"note": "initial"}),
    )


def _build_trade(identifier: str, opened_at: datetime) -> Trade:
    return Trade(
        id=f"trade-{identifier}",
        signal_id=f"sig-{identifier}",
        source_signal_id=f"sig-{identifier}",
        exchange=Exchange.BINANCE,
        symbol="BTCUSDT",
        side=Side.LONG,
        status=TradeStatus.OPENED,
        entry_price=100.0,
        size=1.0,
        timeframe=Timeframe.M15,
        used_margin=100.0,
        opened_at=opened_at,
        created_at=opened_at,
        allow_long=True,
        allow_short=True,
        metadata=TradeMetadata(extra={"note": "initial"}),
    )


def _read_sheet_rows(path: Path, sheet_name: str) -> list[dict[str, object]]:
    workbook = load_workbook(path)
    worksheet = workbook[sheet_name]
    headers = [cell.value for cell in worksheet[1]]
    rows: list[dict[str, object]] = []
    for excel_row in worksheet.iter_rows(min_row=2, values_only=True):
        if all(value is None for value in excel_row):
            continue
        rows.append({header: excel_row[idx] for idx, header in enumerate(headers)})
    return rows


def test_signal_repository_updates_excel(tmp_path) -> None:
    storage, workbook_path = _create_storage(tmp_path)
    repository = SignalRepository(storage)
    now = datetime.now(ZoneInfo("UTC")).replace(microsecond=0)

    signal = _build_signal("1", 10.0, now)
    repository.save(signal)
    loaded_signals = {loaded.id: loaded for loaded in storage.load_signals()}
    assert loaded_signals[signal.id].triggered_at == signal.triggered_at
    assert loaded_signals[signal.id].triggered_at.tzinfo is not None
    assert loaded_signals[signal.id].metrics_snapshot == signal.metrics_snapshot
    assert loaded_signals[signal.id].metadata is not None
    assert "metrics" not in loaded_signals[signal.id].metadata.extra
    assert loaded_signals[signal.id].metadata.extra.get("note") == "initial"

    stored_rows = storage._writer.read_signals()
    assert stored_rows and isinstance(stored_rows[0], StoredSignalRow)
    assert stored_rows[0].score == pytest.approx(signal.score)
    assert stored_rows[0].allow_long is True
    assert stored_rows[0].metrics_snapshot_json is not None
    stored_snapshot = json.loads(stored_rows[0].metrics_snapshot_json)
    assert stored_snapshot["pct_move"] == pytest.approx(signal.metrics_snapshot.pct_move)

    rows = _read_sheet_rows(workbook_path, "Signals")
    assert len(rows) == 1
    assert rows[0]["id"] == signal.id
    assert rows[0]["score"] == signal.score

    updated_signal = replace(signal, score=12.5, updated_at=now + timedelta(minutes=1))
    repository.save(updated_signal)

    rows = _read_sheet_rows(workbook_path, "Signals")
    assert len(rows) == 1
    assert rows[0]["score"] == updated_signal.score
    updated_cell = rows[0]["updated_at"]
    expected_updated = updated_signal.updated_at.isoformat()
    if isinstance(updated_cell, str):
        assert updated_cell.startswith(expected_updated)
    else:
        assert updated_cell == updated_signal.updated_at

    another_signal = _build_signal("2", 8.0, now + timedelta(minutes=5))
    repository.save(another_signal)

    rows = _read_sheet_rows(workbook_path, "Signals")
    assert {row["id"] for row in rows} == {signal.id, another_signal.id}
    scores = {row["id"]: row["score"] for row in rows}
    assert scores[signal.id] == updated_signal.score
    loaded_signals = {loaded.id: loaded for loaded in storage.load_signals()}
    assert loaded_signals[signal.id].updated_at == updated_signal.updated_at
    assert loaded_signals[signal.id].updated_at.tzinfo is not None
    assert loaded_signals[another_signal.id].triggered_at == another_signal.triggered_at
    assert loaded_signals[another_signal.id].triggered_at.tzinfo is not None
    assert loaded_signals[another_signal.id].metrics_snapshot == another_signal.metrics_snapshot


def test_trade_repository_updates_excel(tmp_path) -> None:
    storage, workbook_path = _create_storage(tmp_path)
    repository = TradeRepository(storage)
    now = datetime.now(ZoneInfo("UTC")).replace(microsecond=0)

    trade = _build_trade("1", now)
    repository.save(trade)
    loaded_trades = {loaded.id: loaded for loaded in storage.load_trades()}
    assert loaded_trades[trade.id].opened_at == trade.opened_at
    assert loaded_trades[trade.id].opened_at and loaded_trades[trade.id].opened_at.tzinfo is not None
    assert loaded_trades[trade.id].metadata is not None
    assert loaded_trades[trade.id].metadata.extra.get("note") == "initial"

    rows = _read_sheet_rows(workbook_path, "Trades")
    assert len(rows) == 1
    assert rows[0]["id"] == trade.id
    assert rows[0]["status"] == trade.status.value

    updated_trade = replace(
        trade,
        status=TradeStatus.CLOSED_TP,
        exit_price=105.0,
        tp_price=105.0,
        sl_price=95.0,
        tp_pct=5.0,
        sl_pct=-5.0,
        pnl=5.0,
        pnl_pct=5.0,
        closed_at=now + timedelta(minutes=10),
        updated_at=now + timedelta(minutes=10),
    )
    repository.save(updated_trade)

    rows = _read_sheet_rows(workbook_path, "Trades")
    assert len(rows) == 1
    assert rows[0]["status"] == updated_trade.status.value
    assert rows[0]["exit_price"] == updated_trade.exit_price
    assert rows[0]["pnl"] == updated_trade.pnl
    assert rows[0]["tp_pct"] == updated_trade.tp_pct
    assert rows[0]["sl_pct"] == updated_trade.sl_pct

    second_trade = _build_trade("2", now + timedelta(minutes=1))
    repository.save(second_trade)

    rows = _read_sheet_rows(workbook_path, "Trades")
    assert {row["id"] for row in rows} == {trade.id, second_trade.id}
    status_map = {row["id"]: row["status"] for row in rows}
    assert status_map[trade.id] == updated_trade.status.value
    loaded_trades = {loaded.id: loaded for loaded in storage.load_trades()}
    assert loaded_trades[trade.id].updated_at == updated_trade.updated_at
    assert loaded_trades[trade.id].updated_at and loaded_trades[trade.id].updated_at.tzinfo is not None
    assert loaded_trades[second_trade.id].opened_at == second_trade.opened_at
    assert loaded_trades[second_trade.id].opened_at and loaded_trades[second_trade.id].opened_at.tzinfo is not None

    stored_trades = storage._writer.read_trades()
    assert stored_trades and isinstance(stored_trades[0], StoredTradeRow)
    assert stored_trades[0].entry_price == pytest.approx(trade.entry_price)


def test_state_repository_persists_per_provider(tmp_path) -> None:
    storage, workbook_path = _create_storage(tmp_path)
    default_repo = StateRepository(storage)
    binance_repo = StateRepository(storage, key="binance")
    bybit_repo = StateRepository(storage, key="bybit")

    now = datetime.now(ZoneInfo("UTC")).replace(microsecond=0)
    binance_repo.update_deposit(1_000.0, "USDT", now)
    binance_repo.set_used_amount(250.0)

    later = now + timedelta(minutes=5)
    bybit_repo.update_deposit(2_000.0, "USDT", later)
    bybit_repo.set_used_amount(125.0)

    # Reload repositories to ensure persisted values are scoped.
    reloaded_binance = StateRepository(storage, key="binance")
    reloaded_bybit = StateRepository(storage, key="bybit")

    assert reloaded_binance.get_deposit() == pytest.approx(1_000.0)
    assert reloaded_binance.get_used_amount() == pytest.approx(250.0)
    assert reloaded_binance.get_last_deposit_update() == now
    assert reloaded_binance.get_last_deposit_update() and reloaded_binance.get_last_deposit_update().tzinfo is not None

    assert reloaded_bybit.get_deposit() == pytest.approx(2_000.0)
    assert reloaded_bybit.get_used_amount() == pytest.approx(125.0)
    assert reloaded_bybit.get_last_deposit_update() == later
    assert reloaded_bybit.get_last_deposit_update() and reloaded_bybit.get_last_deposit_update().tzinfo is not None

    # Default scope remains untouched.
    assert default_repo.get_deposit() == 0.0
    assert default_repo.get_used_amount() == 0.0

    rows = _read_sheet_rows(workbook_path, "State")
    keys = {row.get("key") for row in rows}
    assert {"binance", "bybit"}.issubset(keys)

    stored_state = storage._writer.read_state("binance")
    assert stored_state and isinstance(stored_state, StoredStateRow)
    assert stored_state.deposit_amount == pytest.approx(1_000.0)
