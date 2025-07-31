import os
import threading
import time
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta

from openpyxl import Workbook, load_workbook

from domain.models.confidence import Confidence
from domain.models.timeframe import Timeframe
from data.loader import Loader

from config.constants import FILL_OUTDATED_ENTRY, CROSS_MONITOR_HISTORY_BARS
from utils.logger import logw
from utils.setup_log_utils import HEADERS, row_from_signal
from domain.models.setup_log_column import SetupLogColumn

_CONF_ORDER = {
    Confidence.WEAK.name: 0,
    Confidence.MODERATE.name: 1,
    Confidence.STRONG.name: 2,
}

class SetupLog:
    """Класс для работы с журналом сетапов в XLSX-файле."""

    def __init__(self) -> None:
        self.file_path = os.path.join('.generated', 'xls', 'setup_log.xlsx')
        self.lock = threading.Lock()
        self._executor = ThreadPoolExecutor(max_workers=1)
        self._thread: threading.Thread | None = None

    def _ensure_workbook(self):
        os.makedirs(os.path.dirname(self.file_path), exist_ok=True)
        if os.path.exists(self.file_path):
            wb = load_workbook(self.file_path)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.append(HEADERS)
        return wb, ws

    def _write_setup(self, signal, setup_tf: Timeframe) -> None:
        with self.lock:
            wb, ws = self._ensure_workbook()
            new_row = row_from_signal(signal, setup_tf)
            last_idx = None
            for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=False), start=2):
                if row[1].value == signal.symbol:
                    last_idx = idx
            inserted = False
            if last_idx:
                prev_conf = ws.cell(row=last_idx, column=4).value
                if prev_conf in _CONF_ORDER and _CONF_ORDER.get(new_row[3], 0) > _CONF_ORDER.get(prev_conf, 0):
                    for cell in ws[last_idx]:
                        cell.fill = FILL_OUTDATED_ENTRY
                    ws.insert_rows(last_idx + 1)
                    for col, val in enumerate(new_row, start=1):
                        ws.cell(row=last_idx + 1, column=col, value=val)
                    inserted = True
            if not inserted:
                ws.append(new_row)
            wb.save(self.file_path)

    def record_setup(self, signal, setup_tf: Timeframe) -> None:
        """Записать данные сигнала асинхронно."""
        self._executor.submit(self._write_setup, signal, setup_tf)

    def _update_row(self, timestamp: str, symbol: str,
                    main_high_crossed: bool | None = None,
                    correction_low_crossed: bool | None = None,
                    correction_depth_pct: float | None = None,
                    time_to_main_high: int | None = None,
                    outcome: str | None = None) -> None:
        with self.lock:
            wb, ws = self._ensure_workbook()
            headers = [c.value for c in ws[1]]
            try:
                ts_idx = headers.index(SetupLogColumn.TIMESTAMP.value) + 1
                symbol_idx = headers.index(SetupLogColumn.SYMBOL.value) + 1
                mh_idx = headers.index(SetupLogColumn.MAIN_HIGH_CROSSED.value) + 1
                cl_idx = headers.index(SetupLogColumn.CORRECTION_LOW_CROSSED.value) + 1
                cd_idx = headers.index(SetupLogColumn.CORRECTION_DEPTH_PCT.value) + 1
                tm_idx = headers.index(SetupLogColumn.TIME_TO_MAIN_HIGH.value) + 1
                outcome_idx = headers.index(SetupLogColumn.OUTCOME.value) + 1
            except ValueError:
                return
            for row in ws.iter_rows(min_row=2):
                if row[ts_idx - 1].value == timestamp and row[symbol_idx - 1].value == symbol:
                    if main_high_crossed is not None and row[mh_idx - 1].value in (None, ''):
                        row[mh_idx - 1].value = str(bool(main_high_crossed))
                    if correction_low_crossed is not None and row[cl_idx - 1].value in (None, ''):
                        row[cl_idx - 1].value = str(bool(correction_low_crossed))
                    if correction_depth_pct is not None and row[cd_idx - 1].value in (None, ''):
                        row[cd_idx - 1].value = float(correction_depth_pct)
                    if time_to_main_high is not None and row[tm_idx - 1].value in (None, ''):
                        row[tm_idx - 1].value = int(time_to_main_high)
                    if outcome is not None and row[outcome_idx - 1].value in (None, ''):
                        row[outcome_idx - 1].value = str(outcome)
                    break
            wb.save(self.file_path)

    def update_setup_row(self, timestamp: str, symbol: str,
                         main_high_crossed: bool | None = None,
                         correction_low_crossed: bool | None = None,
                         correction_depth_pct: float | None = None,
                         time_to_main_high: int | None = None,
                         outcome: str | None = None) -> None:
        self._executor.submit(
            self._update_row,
            timestamp,
            symbol,
            main_high_crossed,
            correction_low_crossed,
            correction_depth_pct,
            time_to_main_high,
            outcome,
        )

    # noinspection PyBroadException
    def check_crossings(self, loader: Loader) -> None:
        with self.lock:
            if not os.path.exists(self.file_path):
                return
            wb = load_workbook(self.file_path)
            ws = wb.active
            headers = [c.value for c in ws[1]]
            rows = [row for row in ws.iter_rows(min_row=2, values_only=True)]

        try:
            ts_idx = headers.index(SetupLogColumn.TIMESTAMP.value)
            symbol_idx = headers.index(SetupLogColumn.SYMBOL.value)
            sl_idx = headers.index(SetupLogColumn.SL.value)
            tp_idx = headers.index(SetupLogColumn.TP.value)
            tf_idx = headers.index(SetupLogColumn.TF.value)
            depth_idx = headers.index(SetupLogColumn.CORRECTION_DEPTH_PCT.value)
            time_idx = headers.index(SetupLogColumn.TIME_TO_MAIN_HIGH.value)
            outcome_idx = headers.index(SetupLogColumn.OUTCOME.value)
        except ValueError:
            return

        for row in rows:
            ts_str = row[ts_idx]
            symbol = row[symbol_idx]
            if not ts_str or not symbol:
                continue
            outcome_val = row[outcome_idx]
            if outcome_val not in (None, ''):
                continue
            try:
                pump_time = datetime.fromisoformat(ts_str)
            except Exception:
                continue
            tf_str = row[tf_idx]
            try:
                tf = Timeframe(tf_str)
            except Exception:
                tf = Timeframe.M1
            try:
                sl_val = float(row[sl_idx])
                tp_val = float(row[tp_idx])
            except (TypeError, ValueError):
                continue
            pump_range = tp_val - sl_val
            if pump_range <= 0:
                continue
            if datetime.now() - pump_time > timedelta(minutes=CROSS_MONITOR_HISTORY_BARS * tf.minutes):
                continue
            bars = loader.fetch_ohlcvi(symbol, tf, limit=CROSS_MONITOR_HISTORY_BARS)
            pump_level = tp_val + pump_range
            dump_level = tp_val - 0.8 * pump_range
            min_low = float('inf')
            time_to_main = None
            sl_crossed = False
            outcome = None
            for b in bars:
                if b.timestamp < pump_time:
                    continue
                min_low = min(min_low, b.low)
                if b.low <= sl_val:
                    sl_crossed = True
                if time_to_main is None and b.high >= tp_val:
                    time_to_main = int((b.timestamp - pump_time).total_seconds() // 60)
                if b.high >= pump_level:
                    outcome = 'pump'
                    break
                if b.low <= dump_level:
                    outcome = 'dump'
                    break
            if outcome:
                depth_pct = min(80.0, max(0.0, (tp_val - min_low) / pump_range * 100))
                self.update_setup_row(
                    timestamp=ts_str,
                    symbol=symbol,
                    main_high_crossed=time_to_main is not None,
                    correction_low_crossed=sl_crossed,
                    correction_depth_pct=depth_pct,
                    time_to_main_high=time_to_main if outcome == 'pump' else None,
                    outcome=outcome,
                )

    def start_monitor(self, loader: Loader, interval: int = 60) -> None:
        """Запускает фоновую проверку файла на достижение уровней."""

        if self._thread:
            return

        def _loop() -> None:
            while True:
                time.sleep(interval)
                try:
                    self.check_crossings(loader)
                except Exception as error:
                    logw(f"Ошибка проверки лога: {error}")

        self._thread = threading.Thread(target=_loop, daemon=True)
        self._thread.start()

