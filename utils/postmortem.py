
from __future__ import annotations

import csv
import os
import tempfile
import threading
from dataclasses import dataclass
from datetime import datetime
from typing import Iterable, Optional, Tuple

try:
    from openpyxl import Workbook, load_workbook  # type: ignore
    _HAS_OPENPYXL = True
except Exception:
    _HAS_OPENPYXL = False

# --- domain models (import paths match user's project layout) ---
from domain.models.Bar import Bar
from domain.models.Exchange import Exchange
from domain.models.Timeframe import Timeframe


# ========== config ==========

_XLS_PATH = os.getenv("POSTMORTEM_XLS_PATH", ".generated/postmortem/postmortem.xlsx")
_CSV_PATH = os.getenv("POSTMORTEM_CSV_PATH", ".generated/postmortem/postmortem.csv")

_COLS = [
    "ts_logged", "symbol", "exchange", "timeframe", "scenario",
    "level", "sh_last_high",
    "entry_time", "entry_price", "sl_price", "risk_abs", "risk_atr",
    "rr_max", "rr_max_time", "mfe_abs", "tp_price_at_rr_max",
    "sl_hit", "sl_hit_time", "mae_abs",
    "retest_found", "retest_index", "retest_low", "retest_tol",
    "retest_bars_limit", "reaction_bars_limit",
    "spike_found", "spike_index", "spike_high",
    "conditions", "bars_after_confirm",
    # --- virtual TPs and strategies ---
    "tp1_rr", "tp1_hit", "tp1_time",
    "tp2_rr", "tp2_hit", "tp2_time",
    "first_hit",
    "strat_all_tp1_rr",
    "strat_tp1_50_tp2_50_rr",
]


# ========== logging helpers ==========

def log(msg: str) -> None:
    # lightweight console log, user may replace with their logger
    print(msg, flush=True)


def _human_line(
    scenario: str,
    symbol: str, exchange: Exchange, timeframe: Timeframe,
    entry: Optional[float], sl: Optional[float], tp_at_rr: Optional[float],
    rr_max: float, sl_hit: bool, conds: list[str]
) -> str:
    pref = f"{symbol:<12} {exchange.value:<8} {timeframe.value:<4}"
    if entry is not None and sl is not None and tp_at_rr is not None:
        sl_pct = (entry - sl) / entry * 100.0
        tp_pct = (tp_at_rr - entry) / entry * 100.0
        state = "Stop" if sl_hit else ("Take Profit" if rr_max >= 1.0 else "Hold")
        line = f"SL {sl_pct:+.2f}% TP {tp_pct:+.2f}% RR {rr_max:.2f} | {state}"
    else:
        line = "NO TRADE"
    if conds:
        line += " | conds: " + ",".join(conds)
    return f"{pref} {scenario}: {line}"


# ========== sync: process/thread-safe I/O ==========

class _InterProcessFileLock:
    """
    Cross-platform advisory lock using a companion .lock file.
    Uses msvcrt on Windows and fcntl on POSIX where available.
    """
    def __init__(self, path: str):
        self._lock_path = path + ".lock"
        self._fh = None

    def __enter__(self):
        os.makedirs(os.path.dirname(os.path.abspath(self._lock_path)), exist_ok=True)
        self._fh = open(self._lock_path, "a+")
        try:
            if os.name == "nt":
                import msvcrt  # type: ignore
                msvcrt.locking(self._fh.fileno(), msvcrt.LK_LOCK, 1)
            else:
                import fcntl  # type: ignore
                fcntl.flock(self._fh.fileno(), fcntl.LOCK_EX)
        except Exception:
            # best-effort; still hold the file handle reference
            pass
        return self

    def __exit__(self, exc_type, exc, tb):
        try:
            if self._fh:
                if os.name == "nt":
                    try:
                        import msvcrt  # type: ignore
                        self._fh.seek(0)
                        msvcrt.locking(self._fh.fileno(), msvcrt.LK_UNLCK, 1)
                    except Exception:
                        pass
                else:
                    try:
                        import fcntl  # type: ignore
                        fcntl.flock(self._fh.fileno(), fcntl.LOCK_UN)
                    except Exception:
                        pass
        finally:
            try:
                if self._fh:
                    self._fh.close()
            finally:
                self._fh = None


_MUTEX = threading.RLock()


def _atomic_save_wb(wb: "Workbook", path: str) -> None:
    # write to a temp file in the same directory, then replace
    d = os.path.dirname(os.path.abspath(path)) or "."
    os.makedirs(d, exist_ok=True)
    fd, tmp = tempfile.mkstemp(prefix=".~pm_", suffix=".xlsx", dir=d)
    os.close(fd)
    try:
        wb.save(tmp)
        os.replace(tmp, path)
    finally:
        try:
            if os.path.exists(tmp):
                os.remove(tmp)
        except Exception:
            pass


def _ensure_book_locked() -> None:
    if _HAS_OPENPYXL:
        if not os.path.exists(_XLS_PATH):
            wb = Workbook()
            ws = wb.active
            ws.title = "postmortem"
            ws.append(_COLS)
            _atomic_save_wb(wb, _XLS_PATH)
    else:
        if not os.path.exists(_CSV_PATH):
            d = os.path.dirname(os.path.abspath(_CSV_PATH)) or "."
            os.makedirs(d, exist_ok=True)
            with open(_CSV_PATH, "x", encoding="utf-8", newline="") as f:
                csv.writer(f).writerow(_COLS)
                f.flush()
                try:
                    os.fsync(f.fileno())
                except Exception:
                    pass


def _append(values: dict) -> None:
    row = [values.get(k) for k in _COLS]
    lock_target = _XLS_PATH if _HAS_OPENPYXL else _CSV_PATH

    with _MUTEX, _InterProcessFileLock(lock_target):
        _ensure_book_locked()

        if _HAS_OPENPYXL:
            try:
                wb = load_workbook(_XLS_PATH)
                ws = wb.active

                # Upgrade/ensure header
                header = []
                try:
                    for c in ws.iter_rows(min_row=1, max_row=1):
                        header = [cell.value for cell in c]
                        break
                except Exception:
                    header = []
                if header != _COLS:
                    if header and header == _COLS[:len(header)]:
                        for j in range(len(header), len(_COLS)):
                            ws.cell(row=1, column=j+1, value=_COLS[j])
                    else:
                        for j in range(len(_COLS)):
                            ws.cell(row=1, column=j+1, value=_COLS[j])

                ws.append(row)
                _atomic_save_wb(wb, _XLS_PATH)
                return
            except Exception as e:
                log(f"postmortem: xlsx error ({e}); fallback to CSV.")

        # CSV fallback
        with open(_CSV_PATH, "a", encoding="utf-8", newline="") as f:
            csv.writer(f).writerow(row)
            f.flush()
            try:
                os.fsync(f.fileno())
            except Exception:
                pass


# ========== calc helpers ==========

def _atr(b: Bar) -> float:
    return float(b.atr) if b.atr is not None else max(0.0, float(b.high) - float(b.low))


def _first(it: Iterable[int], pred) -> Optional[int]:
    for i in it:
        if pred(i):
            return i
    return None


def _rr_scan_long(bars: list[Bar], start_idx: int, entry: float, sl: float
                 ) -> Tuple[float, Optional[int], float, bool, Optional[int]]:
    """
    Conservative scan: on a bar that touches both SL and TP-side, SL counts first.
    Returns: (rr_max, rr_t, mfe_abs, sl_hit, sl_t)
    """
    risk = entry - sl
    if risk <= 0:
        return 0.0, None, 0.0, False, None

    rr_max = 0.0
    rr_t: Optional[int] = None
    sl_hit = False
    sl_t: Optional[int] = None
    mfe_abs = 0.0

    for i in range(start_idx, len(bars)):
        b = bars[i]
        l = float(b.low)
        h = float(b.high)

        # SL before TP (conservative)
        if l <= sl:
            sl_hit = True
            sl_t = i
            # record mfe within this bar prior to SL (conservative treat as 0)
            break

        # update rr_max / mfe
        rr_here = (h - entry) / risk
        if rr_here > rr_max:
            rr_max = rr_here
            rr_t = i
            mfe_abs = max(mfe_abs, h - entry)

    return rr_max, rr_t, mfe_abs, sl_hit, sl_t


def _rr_scan_short(bars: list[Bar], start_idx: int, entry: float, sl: float
                 ) -> Tuple[float, Optional[int], float, bool, Optional[int]]:
    """
    Conservative scan for shorts: SL counts first if bar hits both sides.
    Returns: (rr_max, rr_t, mfe_abs, sl_hit, sl_t)
    """
    risk = sl - entry
    if risk <= 0:
        return 0.0, None, 0.0, False, None

    rr_max = 0.0
    rr_t: Optional[int] = None
    sl_hit = False
    sl_t: Optional[int] = None
    mfe_abs = 0.0

    for i in range(start_idx, len(bars)):
        b = bars[i]
        l = float(b.low)
        h = float(b.high)

        # SL before TP (conservative)
        if h >= sl:
            sl_hit = True
            sl_t = i
            break

        rr_here = (entry - l) / risk
        if rr_here > rr_max:
            rr_max = rr_here
            rr_t = i
            mfe_abs = max(mfe_abs, entry - l)

    return rr_max, rr_t, mfe_abs, sl_hit, sl_t


def _virt_targets_long(bars: list[Bar], start_idx: int, entry: float, sl: float):
    """Return: (tp1_hit, tp1_t, tp2_hit, tp2_t, first_hit)"""
    risk = entry - sl
    if risk <= 0:
        return False, None, False, None, None
    tp1 = entry + 1.0 * risk
    tp2 = entry + 2.0 * risk

    tp1_hit = False
    tp2_hit = False
    tp1_t: Optional[int] = None
    tp2_t: Optional[int] = None
    first_hit: Optional[str] = None

    for i in range(start_idx, len(bars)):
        b = bars[i]
        l = float(b.low)
        h = float(b.high)

        if l <= sl:
            if first_hit is None:
                first_hit = "SL"
            break
        if (not tp1_hit) and h >= tp1:
            tp1_hit = True
            tp1_t = i
            if first_hit is None:
                first_hit = "TP1"
        if tp1_hit and (not tp2_hit) and h >= tp2:
            tp2_hit = True
            tp2_t = i
            if first_hit is None:
                first_hit = "TP2"
            # keep scanning not needed for our outputs

    return tp1_hit, tp1_t, tp2_hit, tp2_t, first_hit


def _virt_targets_short(bars: list[Bar], start_idx: int, entry: float, sl: float):
    """Return: (tp1_hit, tp1_t, tp2_hit, tp2_t, first_hit)"""
    risk = sl - entry
    if risk <= 0:
        return False, None, False, None, None
    tp1 = entry - 1.0 * risk
    tp2 = entry - 2.0 * risk

    tp1_hit = False
    tp2_hit = False
    tp1_t: Optional[int] = None
    tp2_t: Optional[int] = None
    first_hit: Optional[str] = None

    for i in range(start_idx, len(bars)):
        b = bars[i]
        l = float(b.low)
        h = float(b.high)

        if h >= sl:
            if first_hit is None:
                first_hit = "SL"
            break
        if (not tp1_hit) and l <= tp1:
            tp1_hit = True
            tp1_t = i
            if first_hit is None:
                first_hit = "TP1"
        if tp1_hit and (not tp2_hit) and l <= tp2:
            tp2_hit = True
            tp2_t = i
            if first_hit is None:
                first_hit = "TP2"

    return tp1_hit, tp1_t, tp2_hit, tp2_t, first_hit


# ========== public API ==========

@dataclass(slots=True)
class PostmortemLongRetestInput:
    symbol: str
    exchange: Exchange
    timeframe: Timeframe
    bars_after_confirm: list[Bar]
    level: float
    ret_i: int                 # index of the retest bar inside bars_after_confirm
    reaction_confirm_bars: int
    atr_tol_mult: float = 0.0  # tolerance for "touched" quality check (does not affect fill)


def log_postmortem_long_retest(inp: PostmortemLongRetestInput) -> None:
    bars = inp.bars_after_confirm
    if not bars or inp.ret_i < 0 or inp.ret_i >= len(bars):
        return

    conds: list[str] = ["scenario:long_retest"]
    ret_low = float(bars[inp.ret_i].low)
    atr_here = _atr(bars[inp.ret_i])
    ret_tol = inp.atr_tol_mult * atr_here

    # Quality and fill
    touched = ret_low <= (float(inp.level) + ret_tol)
    if not touched:
        conds.append("no_retest")
        log(_human_line("long_retest", inp.symbol, inp.exchange, inp.timeframe,
                        None, None, None, 0.0, False, conds))
        return

    filled = ret_low <= float(inp.level)
    if not filled:
        conds.append("no_fill")
        _append({
            "ts_logged": datetime.utcnow().isoformat(timespec="seconds"),
            "symbol": inp.symbol, "exchange": inp.exchange.value, "timeframe": inp.timeframe.value,
            "scenario": "long_retest",
            "level": float(inp.level), "sh_last_high": None,
            "entry_time": None, "entry_price": None, "sl_price": None,
            "risk_abs": None, "risk_atr": None,
            "rr_max": 0.0, "rr_max_time": None, "mfe_abs": 0.0, "tp_price_at_rr_max": None,
            "sl_hit": False, "sl_hit_time": None, "mae_abs": 0.0,
            "retest_found": True, "retest_index": inp.ret_i, "retest_low": float(ret_low), "retest_tol": float(ret_tol),
            "retest_bars_limit": None, "reaction_bars_limit": inp.reaction_confirm_bars,
            "spike_found": None, "spike_index": None, "spike_high": None,
            "conditions": ",".join(conds), "bars_after_confirm": len(bars),
            # virtual fields empty
            "tp1_rr": None, "tp1_hit": None, "tp1_time": None,
            "tp2_rr": None, "tp2_hit": None, "tp2_time": None,
            "first_hit": None, "strat_all_tp1_rr": None, "strat_tp1_50_tp2_50_rr": None,
        })
        log(_human_line("long_retest", inp.symbol, inp.exchange, inp.timeframe,
                        None, None, None, 0.0, False, conds))
        return

    # Entry & SL
    entry_idx = inp.ret_i
    entry_time = bars[entry_idx].time
    entry = float(inp.level)               # limit at the level
    sl = float(ret_low)                    # retest low
    risk_abs = entry - sl
    atr_entry = _atr(bars[entry_idx])
    risk_atr = (risk_abs / atr_entry) if atr_entry > 0.0 else None

    if risk_abs <= 0:
        conds.append("invalid_risk")
        _append({
            "ts_logged": datetime.utcnow().isoformat(timespec="seconds"),
            "symbol": inp.symbol, "exchange": inp.exchange.value, "timeframe": inp.timeframe.value,
            "scenario": "long_retest",
            "level": float(inp.level), "sh_last_high": None,
            "entry_time": entry_time.isoformat(), "entry_price": entry, "sl_price": sl,
            "risk_abs": float(risk_abs), "risk_atr": risk_atr,
            "rr_max": 0.0, "rr_max_time": None, "mfe_abs": 0.0, "tp_price_at_rr_max": None,
            "sl_hit": False, "sl_hit_time": None, "mae_abs": 0.0,
            "retest_found": True, "retest_index": inp.ret_i, "retest_low": float(ret_low), "retest_tol": float(ret_tol),
            "retest_bars_limit": None, "reaction_bars_limit": inp.reaction_confirm_bars,
            "spike_found": None, "spike_index": None, "spike_high": None,
            "conditions": ",".join(conds), "bars_after_confirm": len(bars),
            "tp1_rr": None, "tp1_hit": None, "tp1_time": None,
            "tp2_rr": None, "tp2_hit": None, "tp2_time": None,
            "first_hit": None, "strat_all_tp1_rr": None, "strat_tp1_50_tp2_50_rr": None,
        })
        log(_human_line("long_retest", inp.symbol, inp.exchange, inp.timeframe,
                        None, None, None, 0.0, False, conds))
        return

    # RR scan
    rr_max, rr_t, mfe_abs, sl_hit, sl_t = _rr_scan_long(bars, entry_idx, entry, sl)
    rr_time_iso = bars[rr_t].time.isoformat() if rr_t is not None else None
    tp_at_rr = entry + rr_max * risk_abs
    sl_time_iso = bars[sl_t].time.isoformat() if sl_t is not None else None

    # Virtual TPs
    tp1_hit, tp1_t, tp2_hit, tp2_t, first_hit = _virt_targets_long(bars, entry_idx, entry, sl)
    tp1_time_iso = bars[tp1_t].time.isoformat() if tp1_t is not None else None
    tp2_time_iso = bars[tp2_t].time.isoformat() if tp2_t is not None else None

    strat_all_tp1_rr = 1.0 if tp1_hit else -1.0
    if tp2_hit:
        strat_tp1_50_tp2_50_rr = 1.5
    elif tp1_hit and sl_hit:
        strat_tp1_50_tp2_50_rr = 0.0
    elif tp1_hit and not sl_hit:
        strat_tp1_50_tp2_50_rr = 0.5
    else:
        strat_tp1_50_tp2_50_rr = -1.0

    # write row
    _append({
        "ts_logged": datetime.utcnow().isoformat(timespec="seconds"),
        "symbol": inp.symbol, "exchange": inp.exchange.value, "timeframe": inp.timeframe.value,
        "scenario": "long_retest",
        "level": float(inp.level), "sh_last_high": None,
        "entry_time": entry_time.isoformat(), "entry_price": entry, "sl_price": sl,
        "risk_abs": float(risk_abs), "risk_atr": risk_atr,
        "rr_max": float(rr_max), "rr_max_time": rr_time_iso, "mfe_abs": float(mfe_abs), "tp_price_at_rr_max": float(tp_at_rr),
        "sl_hit": bool(sl_hit), "sl_hit_time": sl_time_iso, "mae_abs": float(entry - sl if sl_hit else 0.0),
        "retest_found": True, "retest_index": inp.ret_i, "retest_low": float(ret_low), "retest_tol": float(ret_tol),
        "retest_bars_limit": None, "reaction_bars_limit": inp.reaction_confirm_bars,
        "spike_found": None, "spike_index": None, "spike_high": None,
        "conditions": ",".join(conds), "bars_after_confirm": len(bars),
        "tp1_rr": 1.0, "tp1_hit": bool(tp1_hit), "tp1_time": tp1_time_iso,
        "tp2_rr": 2.0, "tp2_hit": bool(tp2_hit), "tp2_time": tp2_time_iso,
        "first_hit": first_hit,
        "strat_all_tp1_rr": float(strat_all_tp1_rr),
        "strat_tp1_50_tp2_50_rr": float(strat_tp1_50_tp2_50_rr),
    })

    # human log
    msg = _human_line("long_retest", inp.symbol, inp.exchange, inp.timeframe,
                      entry, sl, tp_at_rr, rr_max, sl_hit, conds)
    msg += f" | virt: TP1 {'✔' if tp1_hit else '✖'} TP2 {'✔' if tp2_hit else '✖'}"
    msg += f" | S1R {strat_all_tp1_rr:+.2f}R S50/50 {strat_tp1_50_tp2_50_rr:+.2f}R"
    log(msg)


# ---- Short (basic hook): assume entry at level after two closes below, SL at spike high ----

@dataclass(slots=True)
class PostmortemShortInput:
    symbol: str
    exchange: Exchange
    timeframe: Timeframe
    bars_after_confirm: list[Bar]
    level: float               # level that was reclaimed below
    entry_idx: int             # index where entry is considered
    spike_high: float          # high of stop-hunt spike (SL above)
    reaction_confirm_bars: int = 2
    sh_last_high: Optional[float] = None


def log_postmortem_short_from_sh_last(inp: PostmortemShortInput) -> None:
    bars = inp.bars_after_confirm
    if not bars or inp.entry_idx < 0 or inp.entry_idx >= len(bars):
        return

    conds: list[str] = ["scenario:short_stop_hunt", "close_below_twice"]
    if inp.sh_last_high is not None and inp.spike_high > inp.sh_last_high:
        conds.append("spike_ok")

    entry_idx = inp.entry_idx
    entry_time = bars[entry_idx].time
    entry = float(inp.level)
    sl = float(inp.spike_high)

    risk_abs = sl - entry
    atr_entry = _atr(bars[entry_idx])
    risk_atr = (risk_abs / atr_entry) if atr_entry > 0.0 else None

    if risk_abs <= 0:
        conds.append("invalid_risk")
        _append({
            "ts_logged": datetime.utcnow().isoformat(timespec="seconds"),
            "symbol": inp.symbol, "exchange": inp.exchange.value, "timeframe": inp.timeframe.value,
            "scenario": "short_stop_hunt",
            "level": float(inp.level), "sh_last_high": float(inp.sh_last_high) if inp.sh_last_high is not None else None,
            "entry_time": entry_time.isoformat(), "entry_price": entry, "sl_price": sl,
            "risk_abs": float(risk_abs), "risk_atr": risk_atr,
            "rr_max": 0.0, "rr_max_time": None, "mfe_abs": 0.0, "tp_price_at_rr_max": None,
            "sl_hit": False, "sl_hit_time": None, "mae_abs": 0.0,
            "retest_found": None, "retest_index": None, "retest_low": None, "retest_tol": None,
            "retest_bars_limit": None, "reaction_bars_limit": inp.reaction_confirm_bars,
            "spike_found": True, "spike_index": None, "spike_high": float(inp.spike_high),
            "conditions": ",".join(conds), "bars_after_confirm": len(bars),
            "tp1_rr": None, "tp1_hit": None, "tp1_time": None,
            "tp2_rr": None, "tp2_hit": None, "tp2_time": None,
            "first_hit": None, "strat_all_tp1_rr": None, "strat_tp1_50_tp2_50_rr": None,
        })
        log(_human_line("short_stop_hunt", inp.symbol, inp.exchange, inp.timeframe,
                        None, None, None, 0.0, False, conds))
        return

    rr_max, rr_t, mfe_abs, sl_hit, sl_t = _rr_scan_short(bars, entry_idx, entry, sl)
    rr_time_iso = bars[rr_t].time.isoformat() if rr_t is not None else None
    tp_at_rr = entry - rr_max * risk_abs
    sl_time_iso = bars[sl_t].time.isoformat() if sl_t is not None else None

    tp1_hit, tp1_t, tp2_hit, tp2_t, first_hit = _virt_targets_short(bars, entry_idx, entry, sl)
    tp1_time_iso = bars[tp1_t].time.isoformat() if tp1_t is not None else None
    tp2_time_iso = bars[tp2_t].time.isoformat() if tp2_t is not None else None

    strat_all_tp1_rr = 1.0 if tp1_hit else -1.0
    if tp2_hit:
        strat_tp1_50_tp2_50_rr = 1.5
    elif tp1_hit and sl_hit:
        strat_tp1_50_tp2_50_rr = 0.0
    elif tp1_hit and not sl_hit:
        strat_tp1_50_tp2_50_rr = 0.5
    else:
        strat_tp1_50_tp2_50_rr = -1.0

    _append({
        "ts_logged": datetime.utcnow().isoformat(timespec="seconds"),
        "symbol": inp.symbol, "exchange": inp.exchange.value, "timeframe": inp.timeframe.value,
        "scenario": "short_stop_hunt",
        "level": float(inp.level), "sh_last_high": float(inp.sh_last_high) if inp.sh_last_high is not None else None,
        "entry_time": entry_time.isoformat(), "entry_price": entry, "sl_price": sl,
        "risk_abs": float(risk_abs), "risk_atr": risk_atr,
        "rr_max": float(rr_max), "rr_max_time": rr_time_iso, "mfe_abs": float(mfe_abs), "tp_price_at_rr_max": float(tp_at_rr),
        "sl_hit": bool(sl_hit), "sl_hit_time": sl_time_iso, "mae_abs": float(sl - entry if sl_hit else 0.0),
        "retest_found": None, "retest_index": None, "retest_low": None, "retest_tol": None,
        "retest_bars_limit": None, "reaction_bars_limit": inp.reaction_confirm_bars,
        "spike_found": True, "spike_index": None, "spike_high": float(inp.spike_high),
        "conditions": ",".join(conds), "bars_after_confirm": len(bars),
        "tp1_rr": 1.0, "tp1_hit": bool(tp1_hit), "tp1_time": tp1_time_iso,
        "tp2_rr": 2.0, "tp2_hit": bool(tp2_hit), "tp2_time": tp2_time_iso,
        "first_hit": first_hit,
        "strat_all_tp1_rr": float(strat_all_tp1_rr),
        "strat_tp1_50_tp2_50_rr": float(strat_tp1_50_tp2_50_rr),
    })

    msg = _human_line("short_stop_hunt", inp.symbol, inp.exchange, inp.timeframe,
                      entry, sl, tp_at_rr, rr_max, sl_hit, conds)
    msg += f" | virt: TP1 {'✔' if tp1_hit else '✖'} TP2 {'✔' if tp2_hit else '✖'}"
    msg += f" | S1R {strat_all_tp1_rr:+.2f}R S50/50 {strat_tp1_50_tp2_50_rr:+.2f}R"
    log(msg)
