"""Bootstrap anomalies workbook for live trading sessions."""
from __future__ import annotations

import shutil
from dataclasses import dataclass
from datetime import datetime, timedelta
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import load_workbook

from bot import config
from bot.data.diary import AnomalyRow, WorkbookDiaryBackend
from bot.data.loader import HistoricalRequest, MarketDataLoader
from bot.domain.analyzer import AnalyzerSettings, SignalAnalyzer
from bot.domain.anomaly_optimizer import (
    AnomalySample,
    CandidateEvaluation,
    THRESHOLD_GRID_METADATA,
    ThresholdCandidate,
    load_threshold_candidate,
    optimize_thresholds,
)
from bot.domain.models.anomaly import Anomaly
from bot.domain.models.bar import BarMetrics, BreakDirection
from bot.domain.models.exchange import Exchange
from bot.domain.models.timeframe import Timeframe
from bot.utils.logger import get_logger


@dataclass(frozen=True, slots=True)
class LoggedAnomaly:
    """In-memory representation of an anomaly stored in the workbook."""

    timestamp: datetime
    exchange: Exchange
    symbol: str
    timeframe: Timeframe
    bar_id: str
    open: float
    high: float
    low: float
    close: float
    volume: float
    metrics: BarMetrics


@dataclass(frozen=True, slots=True)
class BootstrapPreparationResult:
    """Artifacts gathered while preparing the live anomalies log."""

    anomalies: list[LoggedAnomaly]
    samples: list[AnomalySample]
    evaluation: CandidateEvaluation | None
    candidate: ThresholdCandidate | None
    settings: AnalyzerSettings | None
    workbook_path: Path


class AnomalyLiveBootstrapper:
    """Ensure the live anomalies log is synchronized with historical data."""

    _BACKTEST_FILENAME = "anomalies.xlsx"
    _LIVE_FILENAME = "anomalies_live.xlsx"
    _ANOMALIES_SHEET_NAME = "anomalies"

    def __init__(
        self,
        *,
        exchange: Exchange,
        loader: MarketDataLoader,
        analyzer: SignalAnalyzer | None = None,
        diary_root: Path | None = None,
    ) -> None:
        self._exchange = exchange
        self._loader = loader
        self._analyzer = analyzer or SignalAnalyzer()
        self._logger = get_logger(__name__)

        base_diary_root = Path(diary_root) if diary_root is not None else Path("var") / "diary"
        if base_diary_root.name == exchange.value:
            exchange_diary_root = base_diary_root
        else:
            exchange_diary_root = base_diary_root / exchange.value

        exchange_diary_root.mkdir(parents=True, exist_ok=True)

        self._base_diary_root = base_diary_root
        self._diary_root = exchange_diary_root
        self._backtest_path = self._diary_root / self._BACKTEST_FILENAME
        self._live_path = self._diary_root / self._LIVE_FILENAME

    def bootstrap(self) -> list[LoggedAnomaly]:
        """Populate the live anomalies workbook and return its contents."""

        self._ensure_live_workbook()
        existing_rows = self._parse_workbook()
        anomalies_to_append = self._backfill_missing(existing_rows)
        if anomalies_to_append:
            backend = WorkbookDiaryBackend(
                self._diary_root, anomalies_filename=self._LIVE_FILENAME
            )
            backend.append_anomalies(self._to_rows(anomalies_to_append))
            backend.flush(["anomalies"])
            self._logger.info(
                "Appended %d anomalies to %s", len(anomalies_to_append), self._live_path
            )
        self._trim_outdated_rows()
        return self._parse_workbook()

    def prepare(self) -> BootstrapPreparationResult:
        """Synchronize the live workbook and derive optimized thresholds."""

        anomalies = self.bootstrap()
        samples = self._to_samples(anomalies)
        evaluation: CandidateEvaluation | None = None
        candidate: ThresholdCandidate | None = None
        settings: AnalyzerSettings | None = None

        if not samples:
            self._logger.info(
                "No anomaly samples available in %s; skipping optimization", self._live_path
            )
            return BootstrapPreparationResult(
                anomalies=anomalies,
                samples=samples,
                evaluation=evaluation,
                candidate=candidate,
                settings=settings,
                workbook_path=self._live_path,
            )

        try:
            base_candidate = load_threshold_candidate(self._live_path)
        except Exception as exc:  # pragma: no cover - defensive logging
            self._logger.warning(
                "Failed to load baseline thresholds from %s: %s",
                self._live_path,
                exc,
            )
            return BootstrapPreparationResult(
                anomalies=anomalies,
                samples=samples,
                evaluation=evaluation,
                candidate=candidate,
                settings=settings,
                workbook_path=self._live_path,
            )

        try:
            grid_deltas = self._compute_grid_deltas(base_candidate)
            evaluation = optimize_thresholds(
                samples,
                base_candidate,
                grid_deltas=grid_deltas,
            )
        except Exception as exc:  # pragma: no cover - defensive logging
            self._logger.warning(
                "Threshold optimization failed using %d samples: %s",
                len(samples),
                exc,
            )
            return BootstrapPreparationResult(
                anomalies=anomalies,
                samples=samples,
                evaluation=evaluation,
                candidate=candidate,
                settings=settings,
                workbook_path=self._live_path,
            )

        candidate = evaluation.candidate
        settings = self._candidate_to_settings(candidate)
        self._logger.info(
            "Optimized thresholds using %d anomaly samples; score %.2f",
            len(samples),
            evaluation.score,
        )
        return BootstrapPreparationResult(
            anomalies=anomalies,
            samples=samples,
            evaluation=evaluation,
            candidate=candidate,
            settings=settings,
            workbook_path=self._live_path,
        )

    def _ensure_live_workbook(self) -> None:
        if self._live_path.exists():
            return
        if self._backtest_path.exists():
            self._live_path.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(self._backtest_path, self._live_path)
            self._logger.info(
                "Initialized live anomalies workbook from %s", self._backtest_path
            )
            return
        self._logger.info(
            "Creating empty live anomalies workbook at %s", self._live_path
        )
        backend = WorkbookDiaryBackend(
            self._diary_root, anomalies_filename=self._LIVE_FILENAME
        )
        backend.flush(["anomalies"])

    def _parse_workbook(self) -> list[LoggedAnomaly]:
        workbook = load_workbook(self._live_path, data_only=True)
        try:
            try:
                sheet = workbook[self._ANOMALIES_SHEET_NAME]
            except KeyError:
                sheet = workbook.active
            header = [cell.value for cell in sheet[1]]
            index = {str(name): position for position, name in enumerate(header)}
            required_columns = [
                "timestamp",
                "exchange",
                "symbol",
                "timeframe",
                "bar_id",
                "open",
                "high",
                "low",
                "close",
                "volume",
                "metrics_pct_move",
                "metrics_relative_volume",
                "metrics_atr_mult",
                "metrics_upper_wick_pct",
                "metrics_body_pct",
                "metrics_lower_wick_pct",
                "metrics_pct_to_low_break",
                "metrics_pct_to_high_break",
                "metrics_break_direction",
            ]
            for column in required_columns:
                if column not in index:
                    raise KeyError(f"Missing '{column}' column in anomalies workbook")
            rows: list[LoggedAnomaly] = []
            for sheet_row in sheet.iter_rows(min_row=2, values_only=True):
                timestamp = self._parse_timestamp(sheet_row[index["timestamp"]])
                if timestamp is None:
                    continue
                exchange_value = sheet_row[index["exchange"]]
                if exchange_value is None:
                    continue
                symbol = sheet_row[index["symbol"]]
                if symbol is None:
                    continue
                timeframe_value = sheet_row[index["timeframe"]]
                if timeframe_value is None:
                    continue
                bar_id = sheet_row[index["bar_id"]]
                if bar_id is None:
                    continue
                open_price = self._float(sheet_row[index["open"]])
                high_price = self._float(sheet_row[index["high"]])
                low_price = self._float(sheet_row[index["low"]])
                close_price = self._float(sheet_row[index["close"]])
                volume = self._float(sheet_row[index["volume"]])
                pct_move = self._float(sheet_row[index["metrics_pct_move"]])
                relative_volume = self._float(
                    sheet_row[index["metrics_relative_volume"]]
                )
                atr_mult = self._float(sheet_row[index["metrics_atr_mult"]])
                upper_wick_pct = self._float(
                    sheet_row[index["metrics_upper_wick_pct"]]
                )
                body_pct = self._float(sheet_row[index["metrics_body_pct"]])
                lower_wick_pct = self._float(
                    sheet_row[index["metrics_lower_wick_pct"]]
                )
                pct_to_low_break = self._float(
                    sheet_row[index["metrics_pct_to_low_break"]]
                )
                pct_to_high_break = self._float(
                    sheet_row[index["metrics_pct_to_high_break"]]
                )
                break_direction_value = sheet_row[index["metrics_break_direction"]]
                if isinstance(exchange_value, Exchange):
                    exchange = exchange_value
                else:
                    try:
                        exchange = Exchange(exchange_value)
                    except ValueError:
                        try:
                            exchange = Exchange(str(exchange_value))
                        except ValueError:
                            continue
                if isinstance(timeframe_value, Timeframe):
                    timeframe = timeframe_value
                else:
                    try:
                        timeframe = Timeframe(str(timeframe_value))
                    except ValueError:
                        continue
                try:
                    break_direction = BreakDirection[str(break_direction_value)]
                except KeyError:
                    break_direction = BreakDirection.NONE
                metrics = BarMetrics(
                    pct_move=pct_move,
                    relative_volume=relative_volume,
                    atr_mult=atr_mult,
                    upper_wick_pct=upper_wick_pct,
                    body_pct=body_pct,
                    lower_wick_pct=lower_wick_pct,
                    pct_to_low_break=pct_to_low_break,
                    pct_to_high_break=pct_to_high_break,
                    break_direction=break_direction,
                )
                rows.append(
                    LoggedAnomaly(
                        timestamp=timestamp,
                        exchange=exchange,
                        symbol=str(symbol),
                        timeframe=timeframe,
                        bar_id=str(bar_id),
                        open=open_price,
                        high=high_price,
                        low=low_price,
                        close=close_price,
                        volume=volume,
                        metrics=metrics,
                    )
                )
            return rows
        finally:
            workbook.close()

    def _backfill_missing(self, rows: Sequence[LoggedAnomaly]) -> list[Anomaly]:
        if not rows:
            return []
        now = datetime.now(tz=config.TIMEZONE)
        grouped: dict[tuple[Exchange, str, Timeframe], list[LoggedAnomaly]] = {}
        for row in rows:
            grouped.setdefault((row.exchange, row.symbol, row.timeframe), []).append(row)
        anomalies: list[Anomaly] = []
        for (exchange, symbol, timeframe), group in grouped.items():
            latest_timestamp = max(row.timestamp for row in group)
            start = latest_timestamp + timedelta(minutes=timeframe.minutes)
            if start >= now:
                continue
            request = HistoricalRequest(
                exchange=exchange,
                symbol=symbol,
                timeframe=timeframe,
                start=start,
                end=now,
                backtest=False,
            )
            for bar in self._loader.load(request):
                if bar.close_time <= latest_timestamp:
                    continue
                _, anomaly = self._analyzer.analyze_bar(bar)
                if anomaly is None:
                    continue
                anomalies.append(anomaly)
                latest_timestamp = max(latest_timestamp, anomaly.timestamp)
        return anomalies

    def _trim_outdated_rows(self) -> None:
        workbook = load_workbook(self._live_path)
        try:
            try:
                sheet = workbook[self._ANOMALIES_SHEET_NAME]
            except KeyError:
                sheet = workbook.active
            header = [cell.value for cell in sheet[1]]
            try:
                timestamp_index = header.index("timestamp") + 1
            except ValueError:
                return
            cutoff = datetime.now(tz=config.TIMEZONE) - config.BACKTEST_MIN_COVERAGE
            removed = 0
            for row_index in range(sheet.max_row, 1, -1):
                cell_value = sheet.cell(row=row_index, column=timestamp_index).value
                timestamp = self._parse_timestamp(cell_value)
                if timestamp is None:
                    continue
                if timestamp < cutoff:
                    sheet.delete_rows(row_index)
                    removed += 1
            workbook.save(self._live_path)
            if removed:
                self._logger.info(
                    "Removed %d outdated anomaly rows older than %s", removed, cutoff
                )
        finally:
            workbook.close()

    @staticmethod
    def _to_rows(anomalies: Iterable[Anomaly]) -> list[AnomalyRow]:
        return [
            AnomalyRow(
                timestamp=anomaly.timestamp,
                exchange=anomaly.exchange,
                symbol=anomaly.symbol,
                timeframe=anomaly.timeframe,
                bar_id=anomaly.bar_id,
                open=anomaly.open,
                high=anomaly.high,
                low=anomaly.low,
                close=anomaly.close,
                volume=anomaly.volume,
                metrics=anomaly.metrics,
                thresholds=anomaly.thresholds,
            )
            for anomaly in anomalies
        ]

    @staticmethod
    def _to_samples(anomalies: Iterable[LoggedAnomaly]) -> list[AnomalySample]:
        return [
            AnomalySample(
                timestamp=anomaly.timestamp,
                exchange=anomaly.exchange,
                symbol=anomaly.symbol,
                timeframe=anomaly.timeframe,
                bar_id=anomaly.bar_id,
                open=anomaly.open,
                high=anomaly.high,
                low=anomaly.low,
                close=anomaly.close,
                volume=anomaly.volume,
                metrics=anomaly.metrics,
            )
            for anomaly in anomalies
        ]

    @staticmethod
    def _candidate_to_settings(candidate: ThresholdCandidate) -> AnalyzerSettings:
        return AnalyzerSettings(
            min_green_move_pct=candidate.min_green_move_pct,
            min_volume_spike=candidate.min_volume_spike,
            min_anomaly_relative_volume=candidate.min_anomaly_relative_volume,
            min_anomaly_atr_mult=candidate.min_anomaly_atr_mult,
            min_anomaly_upper_wick_pct=candidate.min_anomaly_upper_wick_pct,
            min_relative_volume=candidate.min_relative_volume,
            max_relative_volume=candidate.max_relative_volume,
            min_atr_mult=candidate.min_atr_mult,
            min_pct_move=candidate.min_pct_move,
            max_pct_move=candidate.max_pct_move,
            max_upper_wick_pct=candidate.max_upper_wick_pct,
            max_lower_wick_pct=candidate.max_lower_wick_pct,
        )

    @staticmethod
    def _compute_grid_deltas(candidate: ThresholdCandidate) -> dict[str, float]:
        # Baseline filters remain fixed; the optimizer should only explore
        # secondary thresholds derived from the grid deltas below.
        baseline_fields = {
            "min_green_move_pct",
            "min_volume_spike",
            "min_anomaly_relative_volume",
            "min_anomaly_upper_wick_pct",
            "min_anomaly_atr_mult",
        }
        adjustable_fields = [
            "min_relative_volume",
            "max_relative_volume",
            "min_atr_mult",
            "min_pct_move",
            "max_pct_move",
            "max_upper_wick_pct",
            "max_lower_wick_pct",
            "min_rr",
        ]
        deltas: dict[str, float] = {}
        assert not baseline_fields.intersection(adjustable_fields)
        for field in adjustable_fields:
            metadata = THRESHOLD_GRID_METADATA.get(field)
            if metadata is not None and metadata.step > 0:
                deltas[field] = metadata.step
                continue
            baseline = getattr(candidate, field)
            step = abs(baseline) * 0.1
            if step <= 0:
                step = 0.1
            deltas[field] = step
        return deltas

    @staticmethod
    def _parse_timestamp(value: object) -> datetime | None:
        if value is None:
            return None
        if isinstance(value, datetime):
            return value if value.tzinfo else value.replace(tzinfo=config.TIMEZONE)
        if isinstance(value, str):
            try:
                parsed = datetime.fromisoformat(value)
            except ValueError:
                return None
            return parsed if parsed.tzinfo else parsed.replace(tzinfo=config.TIMEZONE)
        return None

    @staticmethod
    def _float(value: object) -> float:
        if value is None:
            return 0.0
        try:
            return float(value)
        except (TypeError, ValueError):
            return 0.0
