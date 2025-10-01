"""Threshold optimization utilities for anomaly backtests."""

from __future__ import annotations

import math
from collections import defaultdict
from collections.abc import Mapping
from dataclasses import dataclass, replace
from datetime import datetime
from pathlib import Path
from random import Random
from typing import DefaultDict, Iterator, Sequence

from openpyxl import load_workbook

from bot import config
from bot.data.diary import WorkbookDiaryBackend
from bot.domain.models.bar import BarMetrics, BreakDirection
from bot.domain.models.exchange import Exchange
from bot.domain.models.timeframe import Timeframe
from bot.utils.datetime_parser import parse_iso_datetime
from bot.utils.logger import get_logger


LOGGER = get_logger(__name__)


@dataclass(frozen=True, slots=True)
class AnomalySample:
    """Parsed anomaly row with the information required for optimization."""

    timestamp: datetime
    exchange: Exchange
    symbol: str
    timeframe: Timeframe
    bar_id: str
    open: float
    high: float
    low: float
    close: float
    volume: float
    metrics: BarMetrics

    @property
    def break_direction(self) -> BreakDirection:
        return self.metrics.break_direction


@dataclass(frozen=True, slots=True)
class ThresholdCandidate:
    """Collection of thresholds used for evaluating workbook anomalies."""

    min_green_move_pct: float
    min_volume_spike: float
    min_anomaly_relative_volume: float
    min_relative_volume: float
    max_relative_volume: float
    min_anomaly_atr_mult: float
    min_atr_mult: float
    min_pct_move: float
    max_pct_move: float
    min_anomaly_upper_wick_pct: float
    max_upper_wick_pct: float
    max_lower_wick_pct: float
    min_rr: float
    initial_deposit: float
    position_fraction: float

    def updated(self, **changes: float) -> "ThresholdCandidate":
        """Return a copy with one or more values replaced."""

        return replace(self, **changes)


@dataclass(frozen=True, slots=True)
class TradeEvaluation:
    """Result of evaluating a single anomaly for one trade direction."""

    filters: dict[str, bool]
    filters_pass: bool
    trade_executed: bool
    rr: float | None
    pnl_pct: float | None
    equity: float


@dataclass(frozen=True, slots=True)
class AnomalyEvaluation:
    """Long/short results for a specific anomaly sample."""

    sample: AnomalySample
    long: TradeEvaluation
    short: TradeEvaluation


@dataclass(frozen=True, slots=True)
class CandidateEvaluation:
    """Aggregated metrics describing a threshold candidate."""

    candidate: ThresholdCandidate
    evaluations: list[AnomalyEvaluation]
    executed_trades: int
    average_trade_return_pct: float
    score: float
    long_final_equity: float
    short_final_equity: float


@dataclass(frozen=True, slots=True)
class GridFieldMetadata:
    """Describe step and bounds applied during grid exploration."""

    step: float
    minimum: float | None = None
    maximum: float | None = None

    def clamp(self, value: float) -> float:
        if self.minimum is not None and value < self.minimum:
            value = self.minimum
        if self.maximum is not None and value > self.maximum:
            value = self.maximum
        return value

    def apply(self, value: float) -> float:
        clamped = self.clamp(value)
        if self.step > 0:
            snapped = round(clamped / self.step) * self.step
        else:
            snapped = clamped
        return self.clamp(snapped)


def _percentage_metadata(
    step: float, *,
    minimum: float | None = 0.0,
    maximum: float | None = 100.0,
) -> GridFieldMetadata:
    """Helper for percentage-based thresholds with common bounds."""

    return GridFieldMetadata(step=step, minimum=minimum, maximum=maximum)


def _bounded_metadata(
    *, step: float, minimum: float, maximum: float
) -> GridFieldMetadata:
    """Helper for thresholds constrained to a numeric range."""

    return GridFieldMetadata(step=step, minimum=minimum, maximum=maximum)


THRESHOLD_GRID_METADATA: dict[str, GridFieldMetadata] = {
    "min_relative_volume": _bounded_metadata(step=1.0, minimum=5.0, maximum=20.0),
    "max_relative_volume": _bounded_metadata(step=10.0, minimum=100.0, maximum=200.0),
    "min_atr_mult": _bounded_metadata(step=0.1, minimum=3.0, maximum=5.0),
    "min_pct_move": _bounded_metadata(step=1.0, minimum=0.0, maximum=20.0),
    "max_pct_move": _bounded_metadata(step=1.0, minimum=10.0, maximum=30.0),
    "max_upper_wick_pct": _percentage_metadata(step=10.0, minimum=50.0, maximum=100.0),
    "max_lower_wick_pct": _percentage_metadata(step=10.0, maximum=100.0),
    "min_rr": _bounded_metadata(step=0.5, minimum=0.0, maximum=3.0),
}

# These fields mirror the baseline-only thresholds in
# ``AnomalyLiveBootstrapper._compute_grid_deltas``. They are not adjusted during
# grid exploration and should retain their original precision when written back
# to the workbook.
BASELINE_THRESHOLD_FIELDS: frozenset[str] = frozenset(
    {
        "min_green_move_pct",
        "min_volume_spike",
        "min_anomaly_relative_volume",
        "min_anomaly_upper_wick_pct",
        "min_anomaly_atr_mult",
    }
)


def _apply_grid_constraints(field: str, value: float) -> float:
    metadata = THRESHOLD_GRID_METADATA.get(field)
    if metadata is None:
        return value
    return metadata.apply(value)


_ANOMALIES_SHEET = "anomalies"
_THRESHOLDS_SHEET = "thresholds"

_THRESHOLD_FIELD_TO_NAME = {
    "min_green_move_pct": "thresholds_min_green_move_pct",
    "min_volume_spike": "thresholds_min_volume_spike",
    "min_anomaly_relative_volume": "thresholds_min_anomaly_relative_volume",
    "min_relative_volume": "thresholds_min_relative_volume",
    "max_relative_volume": "thresholds_max_relative_volume",
    "min_anomaly_atr_mult": "thresholds_min_anomaly_atr_mult",
    "min_atr_mult": "thresholds_min_atr_mult",
    "min_pct_move": "thresholds_min_pct_move",
    "max_pct_move": "thresholds_max_pct_move",
    "min_anomaly_upper_wick_pct": "thresholds_min_anomaly_upper_wick_pct",
    "max_upper_wick_pct": "thresholds_max_upper_wick_pct",
    "max_lower_wick_pct": "thresholds_max_lower_wick_pct",
    "min_rr": "thresholds_min_rr",
    "initial_deposit": "thresholds_initial_deposit",
    "position_fraction": "thresholds_position_fraction",
}


def _build_threshold_field_order() -> tuple[str, ...]:
    reverse_lookup = {
        named_range: field for field, named_range in _THRESHOLD_FIELD_TO_NAME.items()
    }
    order: list[str] = []
    for _, named_range, _ in WorkbookDiaryBackend._ANOMALY_THRESHOLD_LAYOUT:
        field = reverse_lookup.get(named_range)
        if field is not None:
            order.append(field)
    return tuple(order)


THRESHOLD_FIELD_ORDER: tuple[str, ...] = _build_threshold_field_order()


def iter_candidate_thresholds(
    candidate: ThresholdCandidate,
) -> Iterator[tuple[str, float]]:
    """Yield threshold fields and values following the workbook layout order."""

    for field in THRESHOLD_FIELD_ORDER:
        yield field, getattr(candidate, field)


_SKIP_REASON_LABELS = {
    "timestamp_not_datetime": "Timestamp is not a datetime",
    "missing_required_fields": "Missing required fields",
    "invalid_exchange_or_timeframe": "Invalid exchange or timeframe",
}


class ParseDiagnostics:
    """Collect and format diagnostic information while parsing anomalies."""

    __slots__ = ("_entries",)

    def __init__(self) -> None:
        self._entries: DefaultDict[str, list[tuple[int, str | None]]] = defaultdict(list)

    def add(self, reason: str, row_number: int, detail: str | None = None) -> None:
        self._entries[reason].append((row_number, detail))

    @property
    def total_skipped(self) -> int:
        return sum(len(entries) for entries in self._entries.values())

    @property
    def entries(self) -> dict[str, list[tuple[int, str | None]]]:
        return dict(self._entries)

    def summarize(self, *, max_examples: int = 5) -> str:
        """Return a human readable summary of skipped rows."""

        if not self._entries:
            return ""

        parts: list[str] = []
        processed_keys: set[str] = set()

        def _append_reason(reason: str, entries: list[tuple[int, str | None]]) -> None:
            processed_keys.add(reason)
            label = _SKIP_REASON_LABELS.get(reason, reason)
            sample_entries = []
            for row, detail in entries[:max_examples]:
                if detail:
                    sample_entries.append(f"row {row} ({detail})")
                else:
                    sample_entries.append(f"row {row}")
            if len(entries) > max_examples:
                sample_entries.append("...")
            if sample_entries:
                formatted_samples = ", ".join(sample_entries)
                parts.append(f"{label}: {len(entries)} skipped [{formatted_samples}]")
            else:
                parts.append(f"{label}: {len(entries)} skipped")

        for reason in _SKIP_REASON_LABELS:
            entries = self._entries.get(reason)
            if entries:
                _append_reason(reason, entries)

        for reason, entries in self._entries.items():
            if reason not in processed_keys:
                _append_reason(reason, entries)

        return "; ".join(parts)


def parse_anomaly_samples(
    workbook_path: Path, *, diagnostics: ParseDiagnostics | None = None
) -> list[AnomalySample]:
    """Load anomaly samples from a workbook for offline optimization."""

    diagnostics_collector = diagnostics if diagnostics is not None else ParseDiagnostics()
    workbook = load_workbook(workbook_path, data_only=True)
    try:
        try:
            sheet = workbook[_ANOMALIES_SHEET]
        except KeyError:
            sheet = workbook.active
        header = [cell.value for cell in sheet[1]]
        index = {str(name): position for position, name in enumerate(header)}
        required = [
            "timestamp",
            "exchange",
            "symbol",
            "timeframe",
            "bar_id",
            "open",
            "high",
            "low",
            "close",
            "volume",
            "metrics_pct_move",
            "metrics_relative_volume",
            "metrics_atr_mult",
            "metrics_upper_wick_pct",
            "metrics_body_pct",
            "metrics_lower_wick_pct",
            "metrics_pct_to_low_break",
            "metrics_pct_to_high_break",
            "metrics_break_direction",
        ]
        for column in required:
            if column not in index:
                raise KeyError(f"Missing '{column}' column in anomalies workbook")

        samples: list[AnomalySample] = []
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=2, values_only=True), start=2
        ):
            timestamp_value = row[index["timestamp"]]
            timestamp = _parse_timestamp(timestamp_value)
            if timestamp is None:
                diagnostics_collector.add(
                    "timestamp_not_datetime",
                    row_number,
                    f"value={timestamp_value!r}",
                )
                continue
            exchange_value = row[index["exchange"]]
            symbol = row[index["symbol"]]
            timeframe_value = row[index["timeframe"]]
            bar_id = row[index["bar_id"]]
            required_values = {
                "exchange": exchange_value,
                "symbol": symbol,
                "timeframe": timeframe_value,
                "bar_id": bar_id,
            }
            missing_fields = [
                field for field, value in required_values.items() if value in (None, "")
            ]
            if missing_fields:
                diagnostics_collector.add(
                    "missing_required_fields",
                    row_number,
                    f"fields: {', '.join(missing_fields)}",
                )
                continue
            open_price = _float(row[index["open"]])
            high_price = _float(row[index["high"]])
            low_price = _float(row[index["low"]])
            close_price = _float(row[index["close"]])
            volume = _float(row[index["volume"]])
            pct_move = _float(row[index["metrics_pct_move"]])
            relative_volume = _float(row[index["metrics_relative_volume"]])
            atr_mult = _float(row[index["metrics_atr_mult"]])
            upper_wick_pct = _float(row[index["metrics_upper_wick_pct"]])
            body_pct = _float(row[index["metrics_body_pct"]])
            lower_wick_pct = _float(row[index["metrics_lower_wick_pct"]])
            pct_to_low_break = _float(row[index["metrics_pct_to_low_break"]])
            pct_to_high_break = _float(row[index["metrics_pct_to_high_break"]])
            break_direction_value = row[index["metrics_break_direction"]]

            try:
                exchange = (
                    exchange_value
                    if isinstance(exchange_value, Exchange)
                    else Exchange(str(exchange_value))
                )
                timeframe = (
                    timeframe_value
                    if isinstance(timeframe_value, Timeframe)
                    else Timeframe(str(timeframe_value))
                )
            except ValueError as exc:
                diagnostics_collector.add(
                    "invalid_exchange_or_timeframe",
                    row_number,
                    f"values: exchange={exchange_value!r}, timeframe={timeframe_value!r}, error={exc}",
                )
                continue

            try:
                break_direction = BreakDirection[str(break_direction_value)]
            except Exception:
                break_direction = BreakDirection.NONE

            metrics = BarMetrics(
                pct_move=pct_move,
                relative_volume=relative_volume,
                atr_mult=atr_mult,
                upper_wick_pct=upper_wick_pct,
                body_pct=body_pct,
                lower_wick_pct=lower_wick_pct,
                pct_to_low_break=pct_to_low_break,
                pct_to_high_break=pct_to_high_break,
                break_direction=break_direction,
            )

            samples.append(
                AnomalySample(
                    timestamp=timestamp,
                    exchange=exchange,
                    symbol=str(symbol),
                    timeframe=timeframe,
                    bar_id=str(bar_id),
                    open=open_price,
                    high=high_price,
                    low=low_price,
                    close=close_price,
                    volume=volume,
                    metrics=metrics,
                )
            )

        if diagnostics_collector.total_skipped:
            LOGGER.warning(
                "Skipped %d rows while parsing anomalies from %s: %s",
                diagnostics_collector.total_skipped,
                workbook_path,
                diagnostics_collector.summarize(),
            )

        return samples
    finally:
        workbook.close()


def load_threshold_candidate(workbook_path: Path) -> ThresholdCandidate:
    """Read the baseline threshold configuration from the workbook."""

    workbook = load_workbook(workbook_path, data_only=True)
    try:
        if _THRESHOLDS_SHEET not in workbook.sheetnames:
            raise KeyError("Workbook is missing the 'thresholds' sheet")
        sheet = workbook[_THRESHOLDS_SHEET]

        layout = list(WorkbookDiaryBackend._ANOMALY_THRESHOLD_LAYOUT)
        lookup = {name: value for _, name, value in layout}
        layout_index = {name: row_index for row_index, (_, name, _) in enumerate(layout, start=2)}
        values: dict[str, float] = {}
        for field, named_range in _THRESHOLD_FIELD_TO_NAME.items():
            try:
                defined_name = workbook.defined_names[named_range]
                destinations = list(defined_name.destinations)
            except KeyError:
                destinations = []
            cell_value = None
            for sheet_name, cell_address in destinations:
                if sheet_name == sheet.title:
                    cell_value = sheet[cell_address].value
                    break
            if cell_value is None:
                # Fallback to defaults from the layout when no value is set yet.
                if named_range in layout_index:
                    cell_value = sheet.cell(row=layout_index[named_range], column=2).value
            if cell_value is None:
                cell_value = lookup.get(named_range)
            value = float(cell_value)
            values[field] = _apply_grid_constraints(field, value)

        return ThresholdCandidate(**values)
    finally:
        workbook.close()


def evaluate_candidate(
    samples: Sequence[AnomalySample],
    candidate: ThresholdCandidate,
) -> CandidateEvaluation:
    """Evaluate filters and PnL metrics for a threshold candidate."""

    evaluations: list[AnomalyEvaluation] = []
    executed_returns: list[float] = []

    long_equity = candidate.initial_deposit
    short_equity = candidate.initial_deposit

    for sample in samples:
        metrics = sample.metrics

        long_filters = {
            "green_move": metrics.pct_move >= candidate.min_green_move_pct,
            "volume_spike": metrics.relative_volume >= candidate.min_volume_spike,
            "anomaly_relative_volume": metrics.relative_volume >= candidate.min_anomaly_relative_volume,
            "relative_volume_min": metrics.relative_volume >= candidate.min_relative_volume,
            "relative_volume_max": metrics.relative_volume <= candidate.max_relative_volume,
            "anomaly_atr": metrics.atr_mult >= candidate.min_anomaly_atr_mult,
            "atr": metrics.atr_mult > candidate.min_atr_mult,
            "pct_move_min": metrics.pct_move >= candidate.min_pct_move,
            "pct_move_max": metrics.pct_move <= candidate.max_pct_move,
            "upper_wick": metrics.upper_wick_pct < candidate.max_upper_wick_pct,
            "lower_wick": metrics.lower_wick_pct < candidate.max_lower_wick_pct,
        }
        long_rr = _compute_rr(sample.close, sample.high, sample.low, long=True)
        long_filters["rr"] = long_rr is not None and long_rr > candidate.min_rr
        long_filters_pass = all(long_filters.values())
        long_trade_executed = long_filters_pass
        long_pnl_pct = (
            _compute_long_pnl(sample)
            if long_trade_executed and sample.close != 0
            else None
        )
        if long_trade_executed and long_pnl_pct is not None:
            prev_long_equity = long_equity
            long_equity = prev_long_equity * (
                1 + candidate.position_fraction * long_pnl_pct / 100
            )
            profit_delta = (
                prev_long_equity * candidate.position_fraction * long_pnl_pct / 100
            )
            executed_returns.append(profit_delta)
        long_evaluation = TradeEvaluation(
            filters=long_filters,
            filters_pass=long_filters_pass,
            trade_executed=long_trade_executed,
            rr=long_rr,
            pnl_pct=long_pnl_pct,
            equity=long_equity,
        )

        short_filters = {
            "green_move": metrics.pct_move >= candidate.min_green_move_pct,
            "volume_spike": metrics.relative_volume >= candidate.min_volume_spike,
            "anomaly_relative_volume": metrics.relative_volume >= candidate.min_anomaly_relative_volume,
            "relative_volume": (
                metrics.relative_volume < candidate.min_relative_volume
                or metrics.relative_volume > candidate.max_relative_volume
            ),
            "anomaly_atr": metrics.atr_mult >= candidate.min_anomaly_atr_mult,
            "atr": metrics.atr_mult < candidate.min_atr_mult,
            "pct_move": (
                metrics.pct_move > candidate.max_pct_move
                or metrics.pct_move < candidate.min_pct_move
            ),
            "upper_wick": metrics.upper_wick_pct < candidate.max_upper_wick_pct,
            "lower_wick": metrics.lower_wick_pct < candidate.max_lower_wick_pct,
        }
        short_rr = _compute_rr(sample.close, sample.high, sample.low, long=False)
        short_filters["rr"] = short_rr is not None and short_rr >= candidate.min_rr
        short_filters_pass = all(short_filters.values())
        short_trade_executed = short_filters_pass
        short_pnl_pct = (
            _compute_short_pnl(sample)
            if short_trade_executed and sample.close != 0
            else None
        )
        if short_trade_executed and short_pnl_pct is not None:
            prev_short_equity = short_equity
            short_equity = prev_short_equity * (
                1 + candidate.position_fraction * short_pnl_pct / 100
            )
            profit_delta = (
                prev_short_equity * candidate.position_fraction * short_pnl_pct / 100
            )
            executed_returns.append(profit_delta)
        short_evaluation = TradeEvaluation(
            filters=short_filters,
            filters_pass=short_filters_pass,
            trade_executed=short_trade_executed,
            rr=short_rr,
            pnl_pct=short_pnl_pct,
            equity=short_equity,
        )

        evaluations.append(
            AnomalyEvaluation(sample=sample, long=long_evaluation, short=short_evaluation)
        )

    executed_trades = len(executed_returns)
    average_return_pct = sum(executed_returns) / executed_trades if executed_trades else 0.0
    score = average_return_pct * executed_trades if executed_trades else 0.0

    return CandidateEvaluation(
        candidate=candidate,
        evaluations=evaluations,
        executed_trades=executed_trades,
        average_trade_return_pct=average_return_pct,
        score=score,
        long_final_equity=long_equity,
        short_final_equity=short_equity,
    )


def _optimize_iteration(
    samples: Sequence[AnomalySample],
    base_evaluation: CandidateEvaluation,
    grid_deltas: Mapping[str, float],
    random_iterations: int,
    random_scale: float,
    rng: Random,
) -> CandidateEvaluation:
    """Perform a single grid/random exploration step starting from ``base_evaluation``."""

    best_evaluation = base_evaluation
    base_candidate = base_evaluation.candidate

    if not grid_deltas:
        return best_evaluation

    fields = list(grid_deltas.keys())

    def _compute_field_offsets(
        field: str, baseline: float, delta: float
    ) -> list[int]:
        metadata = THRESHOLD_GRID_METADATA.get(field)
        offsets: set[int] = set()

        if metadata and metadata.step > 0:
            step = metadata.step
            if metadata.minimum is not None:
                down_steps = int(
                    math.floor((baseline - metadata.minimum) / step + 1e-9)
                )
                offsets.update(-index for index in range(1, down_steps + 1))
            if metadata.maximum is not None:
                up_steps = int(
                    math.floor((metadata.maximum - baseline) / step + 1e-9)
                )
                offsets.update(index for index in range(1, up_steps + 1))

        if not offsets and delta != 0:
            offsets.update((-1, 1))

        return sorted(offsets)

    def _limited_offsets(offsets: list[int]) -> list[int]:
        if not offsets:
            return [0]
        negatives = [offset for offset in offsets if offset < 0]
        positives = [offset for offset in offsets if offset > 0]
        limited: set[int] = {0}
        if negatives:
            limited.add(min(negatives))
        if positives:
            limited.add(max(positives))
        return sorted(limited)

    field_offsets: dict[str, list[int]] = {}
    pair_offsets: dict[str, list[int]] = {}
    for field in fields:
        baseline = getattr(base_candidate, field)
        delta = grid_deltas[field]
        offsets = _compute_field_offsets(field, baseline, delta)
        field_offsets[field] = offsets
        pair_offsets[field] = _limited_offsets(offsets)

    for field in fields:
        baseline = getattr(base_candidate, field)
        delta = grid_deltas[field]
        for offset in field_offsets[field]:
            adjusted = baseline + delta * offset
            adjusted = _apply_grid_constraints(field, adjusted)
            updates = {
                field: _quantize_threshold(adjusted),
            }
            candidate = base_candidate.updated(**updates)
            evaluation = evaluate_candidate(samples, candidate)
            if evaluation.score > best_evaluation.score:
                best_evaluation = evaluation

    for index, primary_field in enumerate(fields):
        primary_baseline = getattr(base_candidate, primary_field)
        primary_delta = grid_deltas[primary_field]
        for secondary_field in fields[index + 1 :]:
            secondary_baseline = getattr(base_candidate, secondary_field)
            secondary_delta = grid_deltas[secondary_field]
            for primary_offset in pair_offsets[primary_field]:
                for secondary_offset in pair_offsets[secondary_field]:
                    if primary_offset == 0 and secondary_offset == 0:
                        continue
                    primary_adjusted = primary_baseline + primary_delta * primary_offset
                    primary_adjusted = _apply_grid_constraints(
                        primary_field, primary_adjusted
                    )
                    secondary_adjusted = (
                        secondary_baseline + secondary_delta * secondary_offset
                    )
                    secondary_adjusted = _apply_grid_constraints(
                        secondary_field, secondary_adjusted
                    )
                    updates = {
                        primary_field: _quantize_threshold(primary_adjusted),
                        secondary_field: _quantize_threshold(secondary_adjusted),
                    }
                    candidate = base_candidate.updated(**updates)
                    evaluation = evaluate_candidate(samples, candidate)
                    if evaluation.score > best_evaluation.score:
                        best_evaluation = evaluation

    for _ in range(random_iterations):
        updates = {}
        for field, delta in grid_deltas.items():
            baseline = getattr(best_evaluation.candidate, field)
            metadata = THRESHOLD_GRID_METADATA.get(field)
            span: float
            if (
                metadata
                and metadata.minimum is not None
                and metadata.maximum is not None
            ):
                span = (metadata.maximum - metadata.minimum) / 2 * random_scale
            elif metadata and metadata.maximum is not None:
                span = abs(metadata.maximum - baseline) * random_scale
            elif metadata and metadata.minimum is not None:
                span = abs(baseline - metadata.minimum) * random_scale
            else:
                span = abs(delta) * random_scale
            if span == 0:
                value = baseline
            else:
                value = baseline + rng.uniform(-span, span)
            value = _apply_grid_constraints(field, value)
            updates[field] = _quantize_threshold(value)
        candidate = best_evaluation.candidate.updated(**updates)
        evaluation = evaluate_candidate(samples, candidate)
        if evaluation.score > best_evaluation.score:
            best_evaluation = evaluation

    return best_evaluation


def optimize_thresholds(
    samples: Sequence[AnomalySample],
    base_candidate: ThresholdCandidate,
    *,
    grid_deltas: Mapping[str, float],
    random_iterations: int = 50,
    random_scale: float = 0.5,
    rng: Random | None = None,
    improvement_tolerance: float = 0.01,
    min_absolute_improvement: float = 0.01,
) -> CandidateEvaluation:
    """Iteratively explore the search space until improvements plateau."""

    rng = rng or Random()
    delta_map = dict(grid_deltas)

    current_evaluation = evaluate_candidate(samples, base_candidate)
    best_evaluation = current_evaluation
    iterations = 0

    while True:
        iterations += 1
        LOGGER.info(
            "Starting threshold optimization iteration %d (score %.2f).",
            iterations,
            current_evaluation.score,
        )
        iteration_evaluation = _optimize_iteration(
            samples,
            current_evaluation,
            delta_map,
            random_iterations,
            random_scale,
            rng,
        )
        improvement = iteration_evaluation.score - current_evaluation.score
        if improvement <= 0:
            LOGGER.info(
                "Iteration %d did not improve the score (Δ=%.2f); stopping.",
                iterations,
                improvement,
            )
            break

        reference = abs(current_evaluation.score)
        if not math.isclose(reference, 0.0, abs_tol=1e-9):
            relative_improvement = improvement / reference
            LOGGER.info(
                "Iteration %d improved score to %.2f (Δ=%.2f, %+.2f%%).",
                iterations,
                iteration_evaluation.score,
                improvement,
                relative_improvement * 100,
            )
        else:
            relative_improvement = float("inf")
            LOGGER.info(
                "Iteration %d improved score to %.2f (Δ=%.2f).",
                iterations,
                iteration_evaluation.score,
                improvement,
            )

        best_evaluation = iteration_evaluation
        current_evaluation = iteration_evaluation

        if math.isfinite(relative_improvement):
            if relative_improvement < improvement_tolerance:
                LOGGER.info(
                    "Relative improvement %.2f%% below %.2f%% threshold; stopping.",
                    relative_improvement * 100,
                    improvement_tolerance * 100,
                )
                break
        elif improvement < min_absolute_improvement:
            LOGGER.info(
                "Absolute improvement %.2f below %.2f threshold; stopping.",
                improvement,
                min_absolute_improvement,
            )
            break

    LOGGER.info(
        "Completed threshold optimization after %d iteration(s); best score %.2f.",
        iterations,
        best_evaluation.score,
    )
    return best_evaluation


def write_threshold_candidate(workbook_path: Path, candidate: ThresholdCandidate) -> None:
    """Persist the selected candidate back into the workbook."""

    workbook = load_workbook(workbook_path)
    try:
        if _THRESHOLDS_SHEET not in workbook.sheetnames:
            raise KeyError("Workbook is missing the 'thresholds' sheet")
        sheet = workbook[_THRESHOLDS_SHEET]

        layout_index = {
            name: row_index
            for row_index, (_, name, _) in enumerate(
                WorkbookDiaryBackend._ANOMALY_THRESHOLD_LAYOUT, start=2
            )
        }

        for field, named_range in _THRESHOLD_FIELD_TO_NAME.items():
            if named_range not in layout_index:
                continue
            row_index = layout_index[named_range]
            value = getattr(candidate, field)
            if field not in BASELINE_THRESHOLD_FIELDS:
                value = _apply_grid_constraints(field, value)
                value = _quantize_threshold(value)
            sheet.cell(row=row_index, column=2).value = value

        workbook.save(workbook_path)
    finally:
        workbook.close()


def _compute_rr(close: float, high: float, low: float, *, long: bool) -> float | None:
    if close == 0:
        return None
    if long:
        risk = close - low
        reward = high - close
    else:
        risk = high - close
        reward = close - low
    if risk <= 0:
        return None
    return reward / risk


def _compute_long_pnl(sample: AnomalySample) -> float | None:
    close = sample.close
    if close == 0:
        return None
    direction = sample.break_direction
    if direction is BreakDirection.HIGH_FIRST:
        return (sample.high - close) / close * 100
    if direction in (BreakDirection.LOW_FIRST, BreakDirection.BOTH):
        return (sample.low - close) / close * 100
    return 0.0


def _compute_short_pnl(sample: AnomalySample) -> float | None:
    close = sample.close
    if close == 0:
        return None
    direction = sample.break_direction
    if direction is BreakDirection.LOW_FIRST:
        return (close - sample.low) / close * 100
    if direction in (BreakDirection.HIGH_FIRST, BreakDirection.BOTH):
        return (close - sample.high) / close * 100
    return 0.0


def _quantize_threshold(value: float) -> float:
    """Clamp negative values and round to one decimal place."""

    return round(value if value >= 0 else 0.0, 1)


def _parse_timestamp(value: object) -> datetime | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=config.TIMEZONE)
    if isinstance(value, str):
        try:
            return parse_iso_datetime(value, timezone=config.TIMEZONE)
        except ValueError:
            return None
    return None


def _float(value: object) -> float:
    return float(value) if value is not None else 0.0

