"""Workbook writers for signals and trades diaries."""
from __future__ import annotations

import atexit
import signal
import time
import threading
from collections import defaultdict
from logging import Logger
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from queue import Empty, Queue
from threading import Event, Lock, Thread
from types import FrameType
from typing import Callable, Iterable, Literal, Optional, Protocol, Sequence, Union

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.utils.cell import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from bot.domain.models.anomaly import (
    Anomaly,
    AnomalyPerformance,
    AnomalyThresholdSnapshot,
)
from bot.domain.models.bar import BarMetrics
from bot.domain.models.close_reason import CloseReason
from bot.domain.models.exchange import Exchange
from bot.domain.models.signal import Signal, ThresholdSnapshot
from bot.domain.models.signal_direction import SignalDirection
from bot.domain.models.timeframe import Timeframe
from bot.domain.models.trade import Trade
from bot.domain.models.trade_status import TradeStatus
from bot import config
from bot.utils.logger import get_logger


_PERMISSIVE_PERFORMANCE = AnomalyPerformance.permissive()


@dataclass(frozen=True, slots=True)
class Row:
    """Base diary row."""


@dataclass(frozen=True, slots=True)
class SignalRow(Row):
    signal_id: str
    timestamp: datetime
    exchange: Exchange
    symbol: str
    timeframe: Timeframe
    direction: SignalDirection
    entry_price: float
    take_profit_price: float
    stop_loss_price: float
    bar_id: str
    metrics: BarMetrics
    thresholds: ThresholdSnapshot


@dataclass(frozen=True, slots=True)
class TradeRow(Row):
    trade_id: str
    source_signal_id: str
    timestamp_open: datetime
    timestamp_close: Optional[datetime]
    exchange: Exchange
    symbol: str
    timeframe: Timeframe
    side: SignalDirection
    entry_price: float
    take_profit_price: float
    stop_loss_price: float
    requested_qty: float
    executed_qty: float
    status: TradeStatus
    avg_fill_price: Optional[float]
    reason_close: Optional[CloseReason]
    sl_be_at: Optional[datetime]
    order_id: Optional[str]
    stop_order_id: Optional[str]
    take_order_id: Optional[str]
    close_order_id: Optional[str]


@dataclass(frozen=True, slots=True)
class AnomalyRow(Row):
    timestamp: datetime
    exchange: Exchange
    symbol: str
    timeframe: Timeframe
    bar_id: str
    open: float
    high: float
    low: float
    close: float
    volume: float
    metrics: BarMetrics
    thresholds: AnomalyThresholdSnapshot


SheetKey = Literal["signals", "trades", "anomalies"]


class DiaryBackend(Protocol):
    def append_signals(self, rows: Iterable[SignalRow]) -> None:  # pragma: no cover - interface definition
        ...

    def append_trades(self, rows: Iterable[TradeRow]) -> None:  # pragma: no cover - interface definition
        ...

    def append_anomalies(self, rows: Iterable[AnomalyRow]) -> None:  # pragma: no cover - interface definition
        ...

    def append_batch(self, sheet_key: SheetKey, rows: Iterable[Row]) -> None:  # pragma: no cover - interface definition
        ...

    def flush(self, sheet_keys: Optional[Iterable[SheetKey]] = None) -> None:  # pragma: no cover - interface definition
        ...

    def get_anomaly_performance(self) -> AnomalyPerformance:  # pragma: no cover - interface definition
        ...

    def rewrite_anomalies(self, rows: Sequence[AnomalyRow]) -> None:  # pragma: no cover - interface definition
        ...


class WorkbookDiaryBackend(DiaryBackend):
    """Persist diary rows into Excel workbooks under a target directory."""

    _ANOMALY_THRESHOLD_LAYOUT = [
        (
            "Minimum anomaly growth (%)",
            "thresholds_min_green_move_pct",
            config.ANOMALY_MIN_GROWTH_PCT,
        ),
        (
            "Minimum anomaly ATR multiple",
            "thresholds_min_anomaly_atr_mult",
            config.ANOMALY_MIN_ATR_MULT,
        ),
        (
            "Minimum anomaly relative volume",
            "thresholds_min_anomaly_relative_volume",
            config.ANOMALY_MIN_RELATIVE_VOLUME,
        ),
        (
            "Minimum anomaly upper wick (%)",
            "thresholds_min_anomaly_upper_wick_pct",
            config.ANOMALY_MIN_UPPER_WICK_PCT,
        ),
        (
            "Minimum anomaly volume spike",
            "thresholds_min_volume_spike",
            config.ANOMALY_MIN_VOLUME_SPIKE,
        ),
        (
            "Minimum relative volume",
            "thresholds_min_relative_volume",
            config.MIN_REL_VOL,
        ),
        (
            "Maximum relative volume",
            "thresholds_max_relative_volume",
            config.MAX_REL_VOL,
        ),
        ("Minimum ATR multiple", "thresholds_min_atr_mult", config.MIN_ATR_MULT),
        ("Minimum percent move", "thresholds_min_pct_move", config.MIN_PCT_MOVE),
        ("Maximum percent move", "thresholds_max_pct_move", config.MAX_PCT_MOVE),
        ("Initial deposit (USDT)", "thresholds_initial_deposit", 100.0),
        ("Position fraction", "thresholds_position_fraction", 0.1),
        (
            "Maximum upper wick (%)",
            "thresholds_max_upper_wick_pct",
            config.MAX_UPPER_WICK_PCT,
        ),
        (
            "Maximum lower wick (%)",
            "thresholds_max_lower_wick_pct",
            config.MAX_LOWER_WICK_PCT,
        ),
        ("Minimum risk/reward", "thresholds_min_rr", config.MIN_RR),
        (
            "Minimum order size (USDT)",
            "thresholds_min_order_usdt",
            config.MIN_ORDER_USDT,
        ),
        (
            "Order fraction of deposit",
            "thresholds_order_pct_of_deposit",
            config.ORDER_PCT_OF_DEPOSIT,
        ),
    ]

    # Mapping is used by formulas in the ``anomalies`` sheet.
    _ANOMALY_THRESHOLD_CELL_MAP = {
        name: "thresholds!$B$%d" % row_index
        for row_index, (_, name, _) in enumerate(_ANOMALY_THRESHOLD_LAYOUT, start=2)
    }

    _SIGNALS_HEADERS = [
        "signal_id",
        "timestamp",
        "exchange",
        "symbol",
        "timeframe",
        "direction",
        "entry_price",
        "take_profit_price",
        "stop_loss_price",
        "bar_id",
        "metrics_pct_move",
        "metrics_relative_volume",
        "metrics_atr_mult",
        "metrics_upper_wick_pct",
        "metrics_body_pct",
        "metrics_lower_wick_pct",
        "metrics_pct_to_low_break",
        "metrics_pct_to_high_break",
        "metrics_break_direction",
        "thresholds_min_green_move_pct",
        "thresholds_min_volume_spike",
        "thresholds_min_relative_volume",
        "thresholds_max_relative_volume",
        "thresholds_min_atr_mult",
        "thresholds_min_pct_move",
        "thresholds_max_pct_move",
        "thresholds_max_upper_wick_pct",
        "thresholds_max_lower_wick_pct",
    ]

    _TRADES_HEADERS = [
        "trade_id",
        "source_signal_id",
        "timestamp_open",
        "timestamp_close",
        "exchange",
        "symbol",
        "timeframe",
        "side",
        "entry_price",
        "take_profit_price",
        "stop_loss_price",
        "requested_qty",
        "executed_qty",
        "status",
        "avg_fill_price",
        "reason_close",
        "sl_be_at",
        "order_id",
        "stop_order_id",
        "take_order_id",
        "close_order_id",
    ]

    _ANOMALIES_HEADERS = [
        "timestamp",
        "exchange",
        "symbol",
        "timeframe",
        "bar_id",
        "open",
        "high",
        "low",
        "close",
        "volume",
        "long_filters_pass",
        "short_filters_pass",
        "long_trade_executed",
        "short_trade_executed",
        "long_filter_green_move_pass",
        "long_filter_volume_spike_pass",
        "long_filter_anomaly_relative_volume_pass",
        "long_filter_relative_volume_min_pass",
        "long_filter_relative_volume_max_pass",
        "long_filter_anomaly_atr_pass",
        "long_filter_atr_pass",
        "long_filter_pct_move_min_pass",
        "long_filter_pct_move_max_pass",
        "long_filter_upper_wick_pass",
        "long_filter_lower_wick_pass",
        "long_filter_rr_pass",
        "short_filter_green_move_pass",
        "short_filter_volume_spike_pass",
        "short_filter_anomaly_relative_volume_pass",
        "short_filter_relative_volume_pass",
        "short_filter_anomaly_atr_pass",
        "short_filter_atr_pass",
        "short_filter_pct_move_pass",
        "short_filter_upper_wick_pass",
        "short_filter_lower_wick_pass",
        "short_filter_rr_pass",
        "long_rr",
        "short_rr",
        "long_pnl_pct",
        "short_pnl_pct",
        "long_equity_pct",
        "short_equity_pct",
        "metrics_pct_move",
        "metrics_relative_volume",
        "metrics_atr_mult",
        "metrics_upper_wick_pct",
        "metrics_body_pct",
        "metrics_lower_wick_pct",
        "metrics_pct_to_low_break",
        "metrics_pct_to_high_break",
        "metrics_break_direction",
        "thresholds_min_green_move_pct",
        "thresholds_min_volume_spike",
        "thresholds_min_relative_volume",
        "thresholds_min_atr_mult",
        "thresholds_min_anomaly_upper_wick_pct",
    ]

    def __init__(self, base_path: Path, *, anomalies_filename: str = "anomalies.xlsx") -> None:
        self._logger = get_logger(__name__)
        self._base_path = base_path
        self._base_path.mkdir(parents=True, exist_ok=True)

        self._signals_path = self._base_path / "signals.xlsx"
        self._trades_path = self._base_path / "trades.xlsx"
        self._anomalies_path = self._base_path / anomalies_filename

        self._sheet_configs: dict[SheetKey, tuple[Path, Callable[[Row, int], list[object]]]] = {
            "signals": (self._signals_path, self._signal_values),
            "trades": (self._trades_path, self._trade_values),
            "anomalies": (self._anomalies_path, self._anomaly_values),
        }
        self._open_workbooks: dict[SheetKey, tuple[Workbook, Worksheet]] = {}

        self._ensure_workbook(self._signals_path, self._SIGNALS_HEADERS, sheet_name="signals")
        self._ensure_workbook(self._trades_path, self._TRADES_HEADERS, sheet_name="trades")
        self._ensure_workbook(self._anomalies_path, self._ANOMALIES_HEADERS, sheet_name="anomalies")
        workbook = load_workbook(self._anomalies_path)
        try:
            self._ensure_anomaly_threshold_sheet(workbook)
            workbook.save(self._anomalies_path)
        finally:
            workbook.close()

    def append_signals(self, rows: Iterable[SignalRow]) -> None:
        materialized = list(rows)
        if not materialized:
            return
        self.append_batch("signals", materialized)
        self.flush(["signals"])

    def append_trades(self, rows: Iterable[TradeRow]) -> None:
        materialized = list(rows)
        if not materialized:
            return
        self.append_batch("trades", materialized)
        self.flush(["trades"])

    def append_anomalies(self, rows: Iterable[AnomalyRow]) -> None:
        materialized = list(rows)
        if not materialized:
            return
        self.append_batch("anomalies", materialized)
        self.flush(["anomalies"])

    def get_anomaly_performance(self) -> AnomalyPerformance:
        if not self._anomalies_path.exists():
            return _PERMISSIVE_PERFORMANCE

        try:
            workbook = load_workbook(self._anomalies_path, data_only=True)
        except Exception:
            self._logger.exception(
                "Не удалось открыть файл аномалий %s для чтения показателей",
                self._anomalies_path,
            )
            return _PERMISSIVE_PERFORMANCE

        try:
            if self._ANOMALY_THRESHOLD_SHEET_NAME not in workbook.sheetnames:
                self._logger.warning(
                    "В файле %s отсутствует лист '%s' с показателями",
                    self._anomalies_path,
                    self._ANOMALY_THRESHOLD_SHEET_NAME,
                )
                return _PERMISSIVE_PERFORMANCE

            sheet = workbook[self._ANOMALY_THRESHOLD_SHEET_NAME]
            label_to_attr = {
                "Средний результат лонг": "long_profit",
                "Winrate лонг": "long_win_rate",
                "Средний результат шорт": "short_profit",
                "Winrate шорт": "short_win_rate",
            }
            values: dict[str, float] = {}
            for label, value in sheet.iter_rows(min_row=2, max_col=2, values_only=True):
                if label in label_to_attr and isinstance(value, (int, float)):
                    values[label_to_attr[label]] = float(value)

            if not values:
                return _PERMISSIVE_PERFORMANCE

            performance = _PERMISSIVE_PERFORMANCE
            return AnomalyPerformance(
                long_profit=values.get("long_profit", performance.long_profit),
                long_win_rate=values.get("long_win_rate", performance.long_win_rate),
                short_profit=values.get("short_profit", performance.short_profit),
                short_win_rate=values.get("short_win_rate", performance.short_win_rate),
            )
        finally:
            workbook.close()

    def append_batch(self, sheet_key: SheetKey, rows: Iterable[Row]) -> None:
        materialized = list(rows)
        if not materialized:
            return
        path, mapper = self._sheet_configs[sheet_key]
        workbook, sheet = self._ensure_workbook_open(sheet_key, path)
        for row in materialized:
            next_row = sheet.max_row + 1
            values = mapper(row, next_row)
            sheet.append(values)
            appended_row = sheet[next_row]
            for value, cell in zip(values, appended_row):
                if isinstance(value, float):
                    cell.number_format = "0.00"
        self._logger.info(
            "Добавлено %d строк в лист дневника '%s' (%s)",
            len(materialized),
            sheet_key,
            path,
        )

    def flush(self, sheet_keys: Optional[Iterable[SheetKey]] = None) -> None:
        keys = list(sheet_keys) if sheet_keys is not None else list(self._open_workbooks.keys())
        for key in keys:
            workbook_entry = self._open_workbooks.pop(key, None)
            if workbook_entry is None:
                continue
            workbook, _ = workbook_entry
            path, _ = self._sheet_configs[key]
            workbook.save(path)
            workbook.close()
            self._logger.info(
                "Сохранён и закрыт файл дневника '%s' для листа '%s'",
                path,
                key,
            )

    def rewrite_anomalies(self, rows: Sequence[AnomalyRow]) -> None:
        sorted_rows = sorted(rows, key=lambda row: row.timestamp)
        self._logger.info(
            "Перезаписываем %d строк аномалий в отсортированном порядке по времени",
            len(sorted_rows),
        )

        workbook = load_workbook(self._anomalies_path)
        try:
            if "anomalies" in workbook.sheetnames:
                sheet = workbook["anomalies"]
            else:
                sheet = workbook.active

            existing_rows = max(0, sheet.max_row - 1)
            if existing_rows > 0:
                sheet.delete_rows(2, existing_rows)

            next_row_index = sheet.max_row + 1
            for row in sorted_rows:
                values = self._anomaly_values(row, next_row_index)
                sheet.append(values)
                appended_row = sheet[next_row_index]
                for value, cell in zip(values, appended_row):
                    if isinstance(value, float):
                        cell.number_format = "0.00"
                next_row_index += 1

            workbook.save(self._anomalies_path)
        finally:
            workbook.close()

    def _ensure_workbook(self, path: Path, headers: list[str], *, sheet_name: str) -> None:
        if path.exists():
            workbook = load_workbook(path)
            try:
                sheet = workbook.active
                if WorkbookDiaryBackend._needs_header(sheet):
                    sheet.append(headers)
                    workbook.save(path)
                else:
                    existing_header = [cell.value for cell in sheet[1]]
                    header_needs_update = existing_header[: len(headers)] != headers
                    header_needs_update = header_needs_update or len(existing_header) != len(
                        headers
                    )
                    if header_needs_update:
                        for column, header in enumerate(headers, start=1):
                            sheet.cell(row=1, column=column).value = header
                        workbook.save(path)
            finally:
                workbook.close()
            self._logger.info(
                "Открыт существующий файл дневника '%s' для листа '%s'",
                path,
                sheet_name,
            )
            return
        else:
            workbook = Workbook()
            try:
                sheet = workbook.active
                sheet.title = sheet_name
                sheet.append(headers)
                workbook.save(path)
            finally:
                workbook.close()
            self._logger.info(
                "Создан файл дневника '%s' для листа '%s'",
                path,
                sheet_name,
            )

    @staticmethod
    def _needs_header(sheet) -> bool:  # type: ignore[no-any-unimported]
        """Return True if the first row has no values and headers should be added."""

        if sheet.max_row == 0:  # pragma: no cover - openpyxl always has at least 1 row
            return True
        first_row = sheet[1]
        return all(cell.value is None for cell in first_row)

    def _ensure_workbook_open(self, sheet_key: SheetKey, path: Path) -> tuple[Workbook, Worksheet]:
        workbook_entry = self._open_workbooks.get(sheet_key)
        if workbook_entry is not None:
            self._logger.info(
                "Файл для листа '%s' уже открыт по пути '%s'",
                sheet_key,
                path,
            )
            return workbook_entry

        workbook = load_workbook(path)
        sheet = workbook.active
        workbook_entry = (workbook, sheet)
        self._open_workbooks[sheet_key] = workbook_entry
        self._logger.info(
            "Открыт файл дневника '%s' для листа '%s'",
            path,
            sheet_key,
        )
        return workbook_entry

    @staticmethod
    def _format_dt(value: Optional[datetime]) -> Optional[str]:
        if value is None:
            return None
        return value.isoformat()

    def _signal_values(self, row: SignalRow, row_index: int) -> list[object]:
        metrics = row.metrics
        thresholds = row.thresholds
        return [
            row.signal_id,
            self._format_dt(row.timestamp),
            row.exchange.value,
            row.symbol,
            row.timeframe.value,
            row.direction.value,
            row.entry_price,
            row.take_profit_price,
            row.stop_loss_price,
            row.bar_id,
            metrics.pct_move,
            metrics.relative_volume,
            metrics.atr_mult,
            metrics.upper_wick_pct,
            metrics.body_pct,
            metrics.lower_wick_pct,
            metrics.pct_to_low_break,
            metrics.pct_to_high_break,
            metrics.break_direction.name,
            thresholds.min_green_move_pct,
            thresholds.min_volume_spike,
            thresholds.min_relative_volume,
            thresholds.max_relative_volume,
            thresholds.min_atr_mult,
            thresholds.min_pct_move,
            thresholds.max_pct_move,
            thresholds.max_upper_wick_pct,
            thresholds.max_lower_wick_pct,
        ]

    def _trade_values(self, row: TradeRow, row_index: int) -> list[object]:
        return [
            row.trade_id,
            row.source_signal_id,
            self._format_dt(row.timestamp_open),
            self._format_dt(row.timestamp_close),
            row.exchange.value,
            row.symbol,
            row.timeframe.value,
            row.side.value,
            row.entry_price,
            row.take_profit_price,
            row.stop_loss_price,
            row.requested_qty,
            row.executed_qty,
            row.status.value,
            row.avg_fill_price,
            row.reason_close.value if row.reason_close else None,
            self._format_dt(row.sl_be_at),
            row.order_id,
            row.stop_order_id,
            row.take_order_id,
            row.close_order_id,
        ]

    def _anomaly_values(self, row: AnomalyRow, row_index: int) -> list[object]:
        metrics = row.metrics
        thresholds = row.thresholds
        threshold_cells = self._ANOMALY_THRESHOLD_CELL_MAP
        long_rr_formula = (
            f'=IFERROR((G{row_index}-I{row_index})/(I{row_index}-H{row_index}),"")'
        )
        short_rr_formula = (
            f'=IFERROR((I{row_index}-H{row_index})/(G{row_index}-I{row_index}),"")'
        )
        long_filter_green_move_formula = (
            f"=IF($AQ{row_index}>={threshold_cells['thresholds_min_green_move_pct']},TRUE,FALSE)"
        )
        long_filter_volume_spike_formula = (
            f"=IF($AR{row_index}>={threshold_cells['thresholds_min_volume_spike']},TRUE,FALSE)"
        )
        long_filter_anomaly_relative_volume_formula = (
            f"=IF($AR{row_index}>={threshold_cells['thresholds_min_anomaly_relative_volume']},TRUE,FALSE)"
        )
        long_filter_relative_volume_min_formula = (
            f"=IF($AR{row_index}>={threshold_cells['thresholds_min_relative_volume']},TRUE,FALSE)"
        )
        long_filter_relative_volume_max_formula = (
            f"=IF($AR{row_index}<={threshold_cells['thresholds_max_relative_volume']},TRUE,FALSE)"
        )
        long_filter_anomaly_atr_formula = (
            f"=IF($AS{row_index}>={threshold_cells['thresholds_min_anomaly_atr_mult']},TRUE,FALSE)"
        )
        long_filter_atr_formula = (
            f"=IF($AS{row_index}>{threshold_cells['thresholds_min_atr_mult']},TRUE,FALSE)"
        )
        long_filter_pct_move_min_formula = (
            f"=IF($AQ{row_index}>={threshold_cells['thresholds_min_pct_move']},TRUE,FALSE)"
        )
        long_filter_pct_move_max_formula = (
            f"=IF($AQ{row_index}<={threshold_cells['thresholds_max_pct_move']},TRUE,FALSE)"
        )
        long_filter_upper_wick_formula = (
            f"=IF($AT{row_index}<{threshold_cells['thresholds_max_upper_wick_pct']},TRUE,FALSE)"
        )
        long_filter_lower_wick_formula = (
            f"=IF($AV{row_index}<{threshold_cells['thresholds_max_lower_wick_pct']},TRUE,FALSE)"
        )
        long_filter_rr_formula = (
            f"=IF($AK{row_index}>{threshold_cells['thresholds_min_rr']},TRUE,FALSE)"
        )
        short_filter_green_move_formula = (
            f"=IF($AQ{row_index}>={threshold_cells['thresholds_min_green_move_pct']},TRUE,FALSE)"
        )
        short_filter_volume_spike_formula = (
            f"=IF($AR{row_index}>={threshold_cells['thresholds_min_volume_spike']},TRUE,FALSE)"
        )
        short_filter_anomaly_relative_volume_formula = (
            f"=IF($AR{row_index}>={threshold_cells['thresholds_min_anomaly_relative_volume']},TRUE,FALSE)"
        )
        short_filter_relative_volume_formula = (
            f"=IF(OR($AR{row_index}<{threshold_cells['thresholds_min_relative_volume']},"
            f"$AR{row_index}>{threshold_cells['thresholds_max_relative_volume']}),TRUE,FALSE)"
        )
        short_filter_anomaly_atr_formula = (
            f"=IF($AS{row_index}>={threshold_cells['thresholds_min_anomaly_atr_mult']},TRUE,FALSE)"
        )
        short_filter_atr_formula = (
            f"=IF($AS{row_index}<{threshold_cells['thresholds_min_atr_mult']},TRUE,FALSE)"
        )
        short_filter_pct_move_formula = (
            f"=IF(OR($AQ{row_index}>{threshold_cells['thresholds_max_pct_move']},"
            f"$AQ{row_index}<{threshold_cells['thresholds_min_pct_move']}),TRUE,FALSE)"
        )
        short_filter_upper_wick_formula = (
            f"=IF($AT{row_index}<{threshold_cells['thresholds_max_upper_wick_pct']},TRUE,FALSE)"
        )
        short_filter_lower_wick_formula = (
            f"=IF($AV{row_index}<{threshold_cells['thresholds_max_lower_wick_pct']},TRUE,FALSE)"
        )
        short_filter_rr_formula = (
            f"=IF($AL{row_index}>={threshold_cells['thresholds_min_rr']},TRUE,FALSE)"
        )
        long_filters_pass_formula = "=AND(" + ",".join(
            f"{column}{row_index}"
            for column in ("O", "P", "Q", "R", "S", "T", "U", "V", "W", "X", "Y", "Z")
        ) + ")"
        short_filters_pass_formula = "=AND(" + ",".join(
            f"{column}{row_index}"
            for column in ("AA", "AB", "AC", "AD", "AE", "AF", "AG", "AH", "AI", "AJ")
        ) + ")"
        long_trade_executed_formula = f"=IF($K{row_index},TRUE,FALSE)"
        short_trade_executed_formula = f"=IF($L{row_index},TRUE,FALSE)"
        thresholds_min_green_formula = f"={threshold_cells['thresholds_min_green_move_pct']}"
        thresholds_min_volume_formula = f"={threshold_cells['thresholds_min_volume_spike']}"
        thresholds_min_relative_formula = f"={threshold_cells['thresholds_min_relative_volume']}"
        thresholds_min_atr_formula = f"={threshold_cells['thresholds_min_atr_mult']}"
        thresholds_min_anomaly_upper_wick_formula = (
            f"={threshold_cells['thresholds_min_anomaly_upper_wick_pct']}"
        )
        long_pnl_formula = (
            f'=IF($M{row_index},'
            f'IF($I{row_index}=0,"",'
            f'IF($AY{row_index}="HIGH_FIRST",(G{row_index}-I{row_index})/I{row_index}*100,'
            f'IF(OR($AY{row_index}="LOW_FIRST",$AY{row_index}="BOTH"),(H{row_index}-I{row_index})/I{row_index}*100,0))),'
            f'"")'
        )
        short_pnl_formula = (
            f'=IF($N{row_index},'
            f'IF($I{row_index}=0,"",'
            f'IF($AY{row_index}="LOW_FIRST",(I{row_index}-H{row_index})/I{row_index}*100,'
            f'IF(OR($AY{row_index}="HIGH_FIRST",$AY{row_index}="BOTH"),(I{row_index}-G{row_index})/I{row_index}*100,0))),'
            f'"")'
        )
        long_equity_formula = (
            f"=IF(ISNUMBER(AO{row_index-1}),"
            f"IF($M{row_index},AO{row_index-1}*(1+{threshold_cells['thresholds_position_fraction']}*AM{row_index}/100),AO{row_index-1}),"
            f"IF($M{row_index},{threshold_cells['thresholds_initial_deposit']}*(1+{threshold_cells['thresholds_position_fraction']}*AM{row_index}/100),{threshold_cells['thresholds_initial_deposit']}))"
        )
        short_equity_formula = (
            f"=IF(ISNUMBER(AP{row_index-1}),"
            f"IF($N{row_index},AP{row_index-1}*(1+{threshold_cells['thresholds_position_fraction']}*AN{row_index}/100),AP{row_index-1}),"
            f"IF($N{row_index},{threshold_cells['thresholds_initial_deposit']}*(1+{threshold_cells['thresholds_position_fraction']}*AN{row_index}/100),{threshold_cells['thresholds_initial_deposit']}))"
        )
        return [
            self._format_dt(row.timestamp),
            row.exchange.value,
            row.symbol,
            row.timeframe.value,
            row.bar_id,
            row.open,
            row.high,
            row.low,
            row.close,
            row.volume,
            long_filters_pass_formula,
            short_filters_pass_formula,
            long_trade_executed_formula,
            short_trade_executed_formula,
            long_filter_green_move_formula,
            long_filter_volume_spike_formula,
            long_filter_anomaly_relative_volume_formula,
            long_filter_relative_volume_min_formula,
            long_filter_relative_volume_max_formula,
            long_filter_anomaly_atr_formula,
            long_filter_atr_formula,
            long_filter_pct_move_min_formula,
            long_filter_pct_move_max_formula,
            long_filter_upper_wick_formula,
            long_filter_lower_wick_formula,
            long_filter_rr_formula,
            short_filter_green_move_formula,
            short_filter_volume_spike_formula,
            short_filter_anomaly_relative_volume_formula,
            short_filter_relative_volume_formula,
            short_filter_anomaly_atr_formula,
            short_filter_atr_formula,
            short_filter_pct_move_formula,
            short_filter_upper_wick_formula,
            short_filter_lower_wick_formula,
            short_filter_rr_formula,
            long_rr_formula,
            short_rr_formula,
            long_pnl_formula,
            short_pnl_formula,
            long_equity_formula,
            short_equity_formula,
            metrics.pct_move,
            metrics.relative_volume,
            metrics.atr_mult,
            metrics.upper_wick_pct,
            metrics.body_pct,
            metrics.lower_wick_pct,
            metrics.pct_to_low_break,
            metrics.pct_to_high_break,
            metrics.break_direction.name,
            thresholds_min_green_formula,
            thresholds_min_volume_formula,
            thresholds_min_relative_formula,
            thresholds_min_atr_formula,
            thresholds_min_anomaly_upper_wick_formula,
        ]
    @staticmethod
    def _set_named_range(workbook: Workbook, *, name: str, sheet_title: str, column_letter: str, row: int) -> None:
        """Ensure a workbook defined name points at the requested cell."""

        escaped_sheet_title = sheet_title.replace("'", "''")
        attr_text = f"'{escaped_sheet_title}'!${column_letter.upper()}${row}"
        if name in workbook.defined_names:
            del workbook.defined_names[name]
        workbook.defined_names[name] = DefinedName(name=name, attr_text=attr_text)

    _ANOMALY_THRESHOLD_SHEET_NAME = "thresholds"

    def _ensure_anomaly_threshold_sheet(self, workbook: Workbook) -> None:
        """Create the ``thresholds`` sheet with default values when missing."""

        if self._ANOMALY_THRESHOLD_SHEET_NAME in workbook.sheetnames:
            sheet = workbook[self._ANOMALY_THRESHOLD_SHEET_NAME]
        else:
            sheet = workbook.create_sheet(self._ANOMALY_THRESHOLD_SHEET_NAME)

        # Header row for readability.
        if sheet["A1"].value is None:
            sheet["A1"].value = "Parameter"
        if sheet["B1"].value is None:
            sheet["B1"].value = "Value"

        for row_index, (label, name, default_value) in enumerate(
            self._ANOMALY_THRESHOLD_LAYOUT, start=2
        ):
            label_cell = sheet.cell(row=row_index, column=1)
            value_cell = sheet.cell(row=row_index, column=2)

            if label_cell.value is None:
                label_cell.value = label
            if value_cell.value is None:
                value_cell.value = default_value

            self._set_named_range(
                workbook,
                name=name,
                sheet_title=sheet.title,
                column_letter=value_cell.column_letter,
                row=value_cell.row,
            )

        summary_row_start = len(self._ANOMALY_THRESHOLD_LAYOUT) + 2
        long_equity_column_letter = get_column_letter(
            self._ANOMALIES_HEADERS.index("long_equity_pct") + 1
        )
        short_equity_column_letter = get_column_letter(
            self._ANOMALIES_HEADERS.index("short_equity_pct") + 1
        )
        long_pnl_column_letter = get_column_letter(
            self._ANOMALIES_HEADERS.index("long_pnl_pct") + 1
        )
        short_pnl_column_letter = get_column_letter(
            self._ANOMALIES_HEADERS.index("short_pnl_pct") + 1
        )
        long_trade_executed_column_letter = get_column_letter(
            self._ANOMALIES_HEADERS.index("long_trade_executed") + 1
        )
        short_trade_executed_column_letter = get_column_letter(
            self._ANOMALIES_HEADERS.index("short_trade_executed") + 1
        )

        def _column_range(column_letter: str) -> str:
            return f"anomalies!${column_letter.upper()}$2:${column_letter.upper()}$1048576"

        long_equity_range = _column_range(long_equity_column_letter)
        short_equity_range = _column_range(short_equity_column_letter)
        long_pnl_range = _column_range(long_pnl_column_letter)
        short_pnl_range = _column_range(short_pnl_column_letter)
        long_trade_executed_range = _column_range(long_trade_executed_column_letter)
        short_trade_executed_range = _column_range(short_trade_executed_column_letter)

        # The summary formulas below depend on the columns used in the ``anomalies``
        # sheet. We read several columns, so changes to ``_ANOMALIES_HEADERS`` that
        # move these fields must also update the calculated column letters above.
        summary_rows = [
            (
                "Средний результат лонг",
                f"=AVERAGE({long_equity_range})-thresholds_initial_deposit",
            ),
            (
                "Средний результат шорт",
                f"=AVERAGE({short_equity_range})-thresholds_initial_deposit",
            ),
            (
                "Выигрыши лонг",
                f"=COUNTIFS({long_trade_executed_range},TRUE,{long_pnl_range},\">0\")",
            ),
            (
                "Проигрыши лонг",
                f"=COUNTIFS({long_trade_executed_range},TRUE,{long_pnl_range},\"<0\")",
            ),
            (
                "Winrate лонг",
                f"=IFERROR(COUNTIFS({long_trade_executed_range},TRUE,{long_pnl_range},\">0\")/"
                f"(COUNTIFS({long_trade_executed_range},TRUE,{long_pnl_range},\">0\")+"
                f"COUNTIFS({long_trade_executed_range},TRUE,{long_pnl_range},\"<0\")),0)",
            ),
            (
                "Выигрыши шорт",
                f"=COUNTIFS({short_trade_executed_range},TRUE,{short_pnl_range},\">0\")",
            ),
            (
                "Проигрыши шорт",
                f"=COUNTIFS({short_trade_executed_range},TRUE,{short_pnl_range},\"<0\")",
            ),
            (
                "Winrate шорт",
                f"=IFERROR(COUNTIFS({short_trade_executed_range},TRUE,{short_pnl_range},\">0\")/"
                f"(COUNTIFS({short_trade_executed_range},TRUE,{short_pnl_range},\">0\")+"
                f"COUNTIFS({short_trade_executed_range},TRUE,{short_pnl_range},\"<0\")),0)",
            ),
        ]

        for offset, (label, formula) in enumerate(summary_rows, start=0):
            row_index = summary_row_start + offset
            label_cell = sheet.cell(row=row_index, column=1)
            value_cell = sheet.cell(row=row_index, column=2)

            if label_cell.value is None:
                label_cell.value = label
            if value_cell.value is None:
                value_cell.value = formula


@dataclass(slots=True)
class _QueuedRow:
    sheet_key: SheetKey
    row: Row


DiaryRow = Union[SignalRow, TradeRow, AnomalyRow]


@dataclass(slots=True)
class WorkbookDiary:
    path: Path
    backend: Optional[DiaryBackend] = None
    anomalies_filename: str = "anomalies.xlsx"

    _queue: Queue[_QueuedRow] = field(init=False, repr=False)
    _flush_event: Event = field(init=False, repr=False)
    _flush_complete: Event = field(init=False, repr=False)
    _stop_event: Event = field(init=False, repr=False)
    _worker_exception: Optional[BaseException] = field(default=None, init=False, repr=False)
    _worker: Thread = field(init=False, repr=False)
    _batch_size: int = field(init=False, repr=False)
    _flush_timeout: float = field(init=False, repr=False)
    _event_poll_interval: float = field(init=False, repr=False)
    _close_lock: Lock = field(init=False, repr=False)
    _closed: bool = field(default=False, init=False, repr=False)
    _closing: bool = field(default=False, init=False, repr=False)
    _logger: Logger = field(default_factory=lambda: get_logger(__name__), init=False, repr=False)
    _previous_signal_handlers: dict[int, object] = field(init=False, repr=False)
    _signal_handlers: dict[int, Callable[[int, FrameType | None], None]] = field(
        init=False, repr=False
    )
    _anomaly_rows: list[AnomalyRow] = field(init=False, repr=False)

    def __post_init__(self) -> None:
        if self.backend is None:
            self.backend = WorkbookDiaryBackend(
                self.path, anomalies_filename=self.anomalies_filename
            )

        self._batch_size = config.DIARY_BATCH_SIZE
        self._flush_timeout = config.DIARY_FLUSH_TIMEOUT
        self._event_poll_interval = min(0.1, self._flush_timeout)
        self._queue = Queue()
        self._flush_event = Event()
        self._flush_complete = Event()
        self._stop_event = Event()
        self._close_lock = Lock()
        self._worker = Thread(target=self._worker_loop, name="WorkbookDiaryWorker", daemon=False)
        self._worker.start()
        self._previous_signal_handlers = {}
        self._signal_handlers = {}
        self._anomaly_rows = []
        self._register_shutdown_hooks()

    def append_signals(self, signals: Iterable[Signal]) -> None:
        rows = [self._signal_to_row(signal) for signal in signals]
        self._enqueue_rows("signals", rows)

    def append_trades(self, trades: Iterable[Trade]) -> None:
        rows = [self._trade_to_row(trade) for trade in trades]
        self._enqueue_rows("trades", rows)

    def append_anomalies(self, anomalies: Iterable[Anomaly]) -> None:
        rows = [self._anomaly_to_row(anomaly) for anomaly in anomalies]
        self._enqueue_rows("anomalies", rows)

    def get_anomaly_performance(self) -> AnomalyPerformance:
        backend = self.backend
        if backend is None:
            return AnomalyPerformance.permissive()

        try:
            return backend.get_anomaly_performance()
        except Exception:
            self._logger.exception("Ошибка чтения показателей аномалий из дневника")
            return AnomalyPerformance.permissive()

    def flush(self) -> None:
        self._raise_if_worker_failed()
        if self._closed:
            return
        worker = self._worker
        if not worker.is_alive():
            return
        self._queue.join()
        self._flush_complete.clear()
        self._flush_event.set()
        while worker.is_alive():
            if self._flush_complete.wait(timeout=self._flush_timeout):
                break
        self._raise_if_worker_failed()

    def close(self) -> None:
        if self._closed:
            return
        with self._close_lock:
            if self._closed or self._closing:
                return
            self._closing = True
        try:
            worker = self._worker
            try:
                self.flush()
                backend = self.backend
                if backend is not None:
                    try:
                        backend.rewrite_anomalies(self._anomaly_rows)
                    except Exception:
                        self._logger.exception(
                            "Не удалось пересортировать аномалии перед закрытием дневника",
                        )
                self._anomaly_rows.clear()
                self._logger.info("Сброс данных дневника завершён; останавливаем запись")
            finally:
                self._stop_event.set()
                self._flush_event.set()
                if worker.is_alive() and worker is not threading.current_thread():
                    worker.join()
            self._raise_if_worker_failed()
        finally:
            with self._close_lock:
                self._closed = True
                self._closing = False

    def _register_shutdown_hooks(self) -> None:
        atexit.register(self.close)
        for signum in (getattr(signal, "SIGINT", None), getattr(signal, "SIGTERM", None)):
            if signum is None:
                continue
            try:
                previous = signal.getsignal(signum)
            except (OSError, RuntimeError, ValueError):
                continue
            handler = self._make_signal_handler(signum)
            try:
                signal.signal(signum, handler)
            except (OSError, RuntimeError, ValueError):
                continue
            self._previous_signal_handlers[signum] = previous
            self._signal_handlers[signum] = handler

    def _make_signal_handler(
        self, signum: int
    ) -> Callable[[int, FrameType | None], None]:
        def handler(received: int, frame: FrameType | None) -> None:
            try:
                self.close()
            finally:
                previous = self._previous_signal_handlers.get(signum)
                if callable(previous):
                    previous(received, frame)
                else:
                    if previous in (signal.SIG_DFL, None):
                        try:
                            signal.signal(signum, signal.SIG_DFL)
                        except (OSError, RuntimeError, ValueError):
                            pass
                        sigint = getattr(signal, "SIGINT", None)
                        if sigint is not None and received == sigint:
                            raise KeyboardInterrupt
                        if hasattr(signal, "raise_signal"):
                            signal.raise_signal(signum)
                        else:
                            raise SystemExit(0)
                        return
                    if previous is signal.SIG_IGN:
                        return

        return handler

    def _enqueue_rows(self, sheet_key: SheetKey, rows: Iterable[DiaryRow]) -> None:
        materialized_rows = list(rows)
        if not materialized_rows:
            return
        self._raise_if_worker_failed()
        if sheet_key == "anomalies":
            self._anomaly_rows.extend(row for row in materialized_rows if isinstance(row, AnomalyRow))
        for row in materialized_rows:
            self._queue.put(_QueuedRow(sheet_key=sheet_key, row=row))

    def _worker_loop(self) -> None:
        assert self.backend is not None
        backend = self.backend
        buffers: defaultdict[SheetKey, list[Row]] = defaultdict(list)
        pending_backend_flush = False
        poll_interval = max(self._event_poll_interval, 0.01)
        next_flush_deadline = time.monotonic() + self._flush_timeout

        try:
            while True:
                if self._stop_event.is_set():
                    self._drain_queue(buffers)
                    pending_backend_flush |= self._flush_buffers(backend, buffers)
                    if pending_backend_flush:
                        backend.flush()
                        pending_backend_flush = False
                    break

                if self._flush_event.is_set():
                    pending_backend_flush |= self._flush_buffers(backend, buffers)
                    if pending_backend_flush:
                        backend.flush()
                        pending_backend_flush = False
                    self._flush_event.clear()
                    self._flush_complete.set()
                    next_flush_deadline = time.monotonic() + self._flush_timeout
                    continue

                remaining_until_flush = max(0.0, next_flush_deadline - time.monotonic())
                timeout = min(poll_interval, remaining_until_flush) if remaining_until_flush > 0 else 0.0
                try:
                    if timeout == 0.0:
                        queued_row = self._queue.get_nowait()
                    else:
                        queued_row = self._queue.get(timeout=timeout)
                except Empty:
                    if time.monotonic() >= next_flush_deadline:
                        pending_backend_flush |= self._flush_buffers(backend, buffers)
                        if pending_backend_flush:
                            backend.flush()
                            pending_backend_flush = False
                        next_flush_deadline = time.monotonic() + self._flush_timeout
                    continue

                try:
                    buffers[queued_row.sheet_key].append(queued_row.row)
                    buffer = buffers[queued_row.sheet_key]
                    next_flush_deadline = time.monotonic() + self._flush_timeout
                    if len(buffer) >= self._batch_size:
                        backend.append_batch(queued_row.sheet_key, buffer)
                        buffer.clear()
                        pending_backend_flush = True
                finally:
                    self._queue.task_done()

        except BaseException as exc:  # pragma: no cover - defensive programming
            if self._worker_exception is None:
                self._worker_exception = exc
            self._stop_event.set()
            try:
                self._drain_queue(buffers)
                pending_backend_flush |= self._flush_buffers(backend, buffers)
                if pending_backend_flush:
                    backend.flush()
                    pending_backend_flush = False
            except Exception:
                pass
        finally:
            try:
                backend.flush()
            except Exception:
                pass

    def _flush_buffers(
        self, backend: DiaryBackend, buffers: dict[SheetKey, list[Row]]
    ) -> bool:
        flushed = False
        for sheet_key, rows in list(buffers.items()):
            if not rows:
                continue
            backend.append_batch(sheet_key, rows)
            rows.clear()
            flushed = True
        return flushed

    def _drain_queue(self, buffers: dict[SheetKey, list[Row]]) -> None:
        while True:
            try:
                queued_row = self._queue.get_nowait()
            except Empty:
                return
            buffers[queued_row.sheet_key].append(queued_row.row)
            self._queue.task_done()

    def _raise_if_worker_failed(self) -> None:
        if self._worker_exception is not None:
            raise self._worker_exception

    @staticmethod
    def _signal_to_row(signal: Signal) -> SignalRow:
        return SignalRow(
            signal_id=signal.signal_id,
            timestamp=signal.timestamp,
            exchange=signal.exchange,
            symbol=signal.symbol,
            timeframe=signal.timeframe,
            direction=signal.direction,
            entry_price=signal.levels.entry_price,
            take_profit_price=signal.levels.take_profit_price,
            stop_loss_price=signal.levels.stop_loss_price,
            bar_id=signal.bar.bar_id,
            metrics=signal.bar.metrics,
            thresholds=signal.thresholds,
        )

    @staticmethod
    def _trade_to_row(trade: Trade) -> TradeRow:
        return TradeRow(
            trade_id=trade.trade_id,
            source_signal_id=trade.source_signal_id,
            timestamp_open=trade.timestamp_open,
            timestamp_close=trade.timestamp_close,
            exchange=trade.exchange,
            symbol=trade.symbol,
            timeframe=trade.timeframe,
            side=trade.side,
            entry_price=trade.entry_price,
            take_profit_price=trade.take_profit_price,
            stop_loss_price=trade.stop_loss_price,
            requested_qty=trade.requested_qty,
            executed_qty=trade.executed_qty,
            status=trade.status,
            avg_fill_price=trade.avg_fill_price,
            reason_close=trade.reason_close,
            sl_be_at=trade.sl_be_at,
            order_id=trade.order_id,
            stop_order_id=trade.stop_order_id,
            take_order_id=trade.take_order_id,
            close_order_id=trade.close_order_id,
        )

    @staticmethod
    def _anomaly_to_row(anomaly: Anomaly) -> AnomalyRow:
        return AnomalyRow(
            timestamp=anomaly.timestamp,
            exchange=anomaly.exchange,
            symbol=anomaly.symbol,
            timeframe=anomaly.timeframe,
            bar_id=anomaly.bar_id,
            open=anomaly.open,
            high=anomaly.high,
            low=anomaly.low,
            close=anomaly.close,
            volume=anomaly.volume,
            metrics=anomaly.metrics,
            thresholds=anomaly.thresholds,
        )
