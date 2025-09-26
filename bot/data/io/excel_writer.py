from __future__ import annotations
from dataclasses import dataclass
from pathlib import Path
from threading import RLock
from typing import Any, Iterable, List, Optional, Sequence, cast

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from .excel_rows import SignalPayload, SignalRow, StateRow, TradePayload, TradeRow


@dataclass(frozen=True)
class SheetConfig:
    name: str
    headers: List[str]


class ExcelWriter:
    """Utility to persist trading data to an Excel workbook."""

    SIGNAL_SHEET = SheetConfig(
        name="Signals",
        headers=[
            "id",
            "candle_id",
            "symbol",
            "exchange",
            "timeframe",
            "candle_timeframe",
            "side",
            "direction",
            "score",
            "triggered_at",
            "created_at",
            "updated_at",
            "allow_long",
            "allow_short",
            "candle_open",
            "candle_high",
            "candle_low",
            "candle_close",
            "candle_volume",
            "candle_quote_volume",
            "candle_started_at",
            "candle_closed_at",
            "thresholds_json",
            "metrics_json",
            "metrics_snapshot_json",
            "metadata_json",
        ],
    )

    TRADE_SHEET = SheetConfig(
        name="Trades",
        headers=[
            "id",
            "signal_id",
            "source_signal_id",
            "exchange",
            "symbol",
            "timeframe",
            "side",
            "status",
            "entry_price",
            "size",
            "used_margin",
            "exit_price",
            "tp_price",
            "sl_price",
            "tp_pct",
            "sl_pct",
            "opened_at",
            "closed_at",
            "pnl",
            "pnl_pct",
            "created_at",
            "updated_at",
            "allow_long",
            "allow_short",
            "thresholds_snapshot_json",
            "metadata_json",
        ],
    )

    STATE_SHEET = SheetConfig(
        name="State",
        headers=["key", "asset", "deposit_amount", "deposit_updated_at", "used_amount"],
    )

    def __init__(self, path: Path) -> None:
        self._path = path
        self._lock = RLock()
        self._initialize_workbook()

    # ------------------------------------------------------------------
    # Public API
    # ------------------------------------------------------------------
    def write_signal(self, row: SignalRow) -> None:
        with self._lock:
            wb = self._load()
            ws = self._get_sheet(wb, self.SIGNAL_SHEET)
            self._upsert_row(ws, self.SIGNAL_SHEET.headers, "id", row.as_excel_values())
            wb.save(self._path)

    def read_signals(self) -> List[SignalPayload]:
        with self._lock:
            wb = self._load()
            ws = self._get_sheet(wb, self.SIGNAL_SHEET)
            return [
                row.to_payload()
                for row in self._iter_rows(ws, self.SIGNAL_SHEET.headers, SignalRow)
            ]

    def write_trade(self, row: TradeRow) -> None:
        with self._lock:
            wb = self._load()
            ws = self._get_sheet(wb, self.TRADE_SHEET)
            self._upsert_row(ws, self.TRADE_SHEET.headers, "id", row.as_excel_values())
            wb.save(self._path)

    def read_trades(self) -> List[TradePayload]:
        with self._lock:
            wb = self._load()
            ws = self._get_sheet(wb, self.TRADE_SHEET)
            return [
                row.to_payload()
                for row in self._iter_rows(ws, self.TRADE_SHEET.headers, TradeRow)
            ]

    def write_state(self, row: StateRow) -> None:
        with self._lock:
            wb = self._load()
            ws = self._get_sheet(wb, self.STATE_SHEET)
            self._upsert_row(ws, self.STATE_SHEET.headers, "key", row.as_excel_values())
            wb.save(self._path)

    def read_state(self, key: Optional[str] = None) -> Optional[StateRow]:
        with self._lock:
            wb = self._load()
            ws = self._get_sheet(wb, self.STATE_SHEET)
            if key is None:
                for row in self._iter_rows(ws, self.STATE_SHEET.headers, StateRow):
                    return row
            else:
                for row in self._iter_rows(ws, self.STATE_SHEET.headers, StateRow):
                    if row.key == key:
                        return row
            return None

    # ------------------------------------------------------------------
    # Internal helpers
    # ------------------------------------------------------------------
    def _initialize_workbook(self) -> None:
        if self._path.exists():
            return
        self._path.parent.mkdir(parents=True, exist_ok=True)
        wb = Workbook()
        # Replace the default sheet with signals to keep order predictable.
        default = wb.active
        if default is None:
            default = wb.create_sheet(self.SIGNAL_SHEET.name)
        else:
            default.title = self.SIGNAL_SHEET.name
        self._write_headers(default, self.SIGNAL_SHEET.headers)
        wb.create_sheet(self.TRADE_SHEET.name)
        wb.create_sheet(self.STATE_SHEET.name)
        self._write_headers(wb[self.TRADE_SHEET.name], self.TRADE_SHEET.headers)
        self._write_headers(wb[self.STATE_SHEET.name], self.STATE_SHEET.headers)
        wb.save(self._path)

    def _load(self) -> Workbook:
        return load_workbook(self._path)

    def _get_sheet(self, workbook: Workbook, config: SheetConfig) -> Worksheet:
        if config.name not in workbook.sheetnames:
            ws = cast(Worksheet, workbook.create_sheet(config.name))
            self._write_headers(ws, config.headers)
            return ws
        ws = cast(Worksheet, workbook[config.name])
        if ws.max_row == 0:
            self._write_headers(ws, config.headers)
        if config is self.STATE_SHEET:
            self._upgrade_state_sheet(ws)
        return ws

    def _upgrade_state_sheet(self, sheet: Worksheet) -> None:
        expected = self.STATE_SHEET.headers
        existing_headers = [cell.value for cell in sheet[1]] if sheet.max_row else []
        if existing_headers[: len(expected)] == list(expected):
            return
        rows: List[StateRow] = []
        if existing_headers:
            normalized_headers = [
                str(header) if header is not None else ""
                for header in existing_headers
            ]
            for values in sheet.iter_rows(min_row=2, values_only=True):
                if all(value is None for value in values):
                    continue
                rows.append(StateRow.from_excel_row(normalized_headers, values))
        self._write_headers(sheet, expected)
        for row in rows:
            sheet.append(row.as_excel_values())

    def _write_headers(self, sheet: Worksheet, headers: Iterable[str]) -> None:
        sheet.delete_rows(1, sheet.max_row)
        sheet.append(list(headers))

    def _iter_rows(
        self, sheet: Worksheet, headers: List[str], row_type
    ) -> Iterable:
        for row in sheet.iter_rows(min_row=2, values_only=True):
            if all(value is None for value in row):
                continue
            yield row_type.from_excel_row(headers, row)

    def _upsert_row(
        self,
        sheet: Worksheet,
        headers: List[str],
        key_column: str,
        row_values: Sequence[object],
    ) -> None:
        row_list = list(row_values)
        if len(row_list) < len(headers):
            row_list.extend([None] * (len(headers) - len(row_list)))
        key_index = headers.index(key_column)
        key_value = row_list[key_index]
        if key_value is None:
            raise ValueError(f"Missing '{key_column}' in row data")
        for excel_row in sheet.iter_rows(min_row=2):
            cell_value = excel_row[key_index].value
            if cell_value == key_value:
                for idx, _ in enumerate(headers):
                    excel_row[idx].value = cast(Any, row_list[idx])
                break
        else:
            sheet.append(row_list)

